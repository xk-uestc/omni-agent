"""Run and score the authored 25-turn multi-source dialogue evaluation."""
from __future__ import annotations

from datetime import datetime, timezone
import hashlib
import json
import math
import os
from pathlib import Path
import re
import sqlite3
import statistics
import time
import uuid

import requests


ROOT = Path(__file__).resolve().parents[1]
QUESTIONS = ROOT / "benchmarks/authored_multiturn_20261008/questions.json"
DATABASE = ROOT / "ict-track8/data/demo_sales.sqlite"
BASE_URL = os.environ.get("ICT8_EVAL_BASE_URL", "http://127.0.0.1:8031").rstrip('/')
SOURCE_FILES = (
    ROOT / "samples/documents/forecast-report.pdf",
    ROOT / "samples/documents/region-targets.xlsx",
    ROOT / "samples/documents/policy-history.md",
    ROOT / "samples/documents/return-policy.docx",
    ROOT / "samples/documents/champion-method.md",
)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def rows_equal(actual, expected, tolerance=0.011) -> bool:
    if len(actual) != len(expected):
        return False
    for actual_row, expected_row in zip(actual, expected):
        if len(actual_row) != len(expected_row):
            return False
        for left, right in zip(actual_row, expected_row):
            if isinstance(right, (int, float)) and not isinstance(right, bool):
                if not isinstance(left, (int, float)) or not math.isclose(float(left), float(right), abs_tol=tolerance):
                    return False
            elif left != right:
                return False
    return True


def response_rows(result):
    columns = result.get("columns", [])
    rows = result.get("rows", [])
    if not isinstance(rows, list):
        return []
    if columns and all(isinstance(row, dict) for row in rows):
        return [[row.get(column) for column in columns if column != '排名'] for row in rows]
    return [list(row.values()) if isinstance(row, dict) else row for row in rows]


def walk(value, *, path=()):
    if isinstance(value, dict):
        yield path, value
        for key, child in value.items():
            yield from walk(child, path=path + (key,))
    elif isinstance(value, list):
        for index, child in enumerate(value):
            yield from walk(child, path=path + (index,))
    else:
        yield path, value


def evidence_text(value):
    pieces = []
    ignored = {"question", "sql", "parameters", "source_uri", "locator", "sha256", "query_hash"}
    for path, item in walk(value):
        if path and path[-1] in ignored:
            continue
        if isinstance(item, str):
            pieces.append(item)
    return "\n".join(pieces)


def answer_text(data):
    result = data.get("result", {}) if isinstance(data, dict) else {}
    answer = result.get("answer", "") if isinstance(result, dict) else ""
    return answer if isinstance(answer, str) else json.dumps(answer, ensure_ascii=False)


def document_ids(data):
    found = set()
    for _, item in walk(data):
        if not isinstance(item, dict):
            continue
        metadata = item.get("metadata")
        if isinstance(metadata, dict) and metadata.get("document_id"):
            found.add(str(metadata["document_id"]))
        if item.get("document_id"):
            found.add(str(item["document_id"]))
    return sorted(found)


def numeric_values(value):
    output = []
    excluded = {"question", "sql", "parameters", "source_uri", "locator", "snippet", "text", "raw_page_text",
        "plan", "trace", "execution_plan", "user_constraint_validation", "generation_attempts", "source_absence_audit"}
    for path, item in walk(value):
        if any(part in excluded for part in path if isinstance(part,str)):
            continue
        if isinstance(item, bool):
            continue
        if isinstance(item, (int, float)):
            output.append(float(item))
        elif isinstance(item, str):
            for raw in re.findall(r"(?<![A-Za-z0-9_])(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?", item):
                try:
                    output.append(float(raw.replace(",", "")))
                except ValueError:
                    pass
    return output


def value_present(data, expected, tolerance=0.02):
    return any(math.isclose(value, float(expected), abs_tol=tolerance) for value in numeric_values(data))


def grade(case, data, expected_rows, source_ids, turn_index):
    result = data.get("result", {}) if isinstance(data, dict) else {}
    answer = answer_text(data)
    material = evidence_text(result)
    checks = {
        "http_status_ok": data.get("http_status") == 200,
        "agent_status_ok": data.get("status") == "ok",
    }
    kind = case["kind"]
    if kind.startswith("sql"):
        checks["sql_route"] = data.get("route") == "sql"
        actual_rows=response_rows(result)
        if re.search(r'\d{4}年?比\d{4}年?(?:增加|增长|多|变化)(?:了)?多少',case['question']) and '变化量' in result.get('columns',[]):
            actual_rows=[[row['变化量']] for row in result['rows']]
        checks["independent_gold_rows_match"] = rows_equal(actual_rows, expected_rows)
    if kind.startswith("evidence"):
        checks["document_or_fusion_route"] = data.get("route") in {"document", "fusion"}
    if kind.startswith("formula"):
        checks["fusion_route"] = data.get("route") == "fusion"
    if kind in {"sql_fact", "evidence_fact"}:
        checks["independent_fact_present"] = all(term in material for term in case.get("required_terms", []))
        if case.get("required_numbers_from_sql"):
            checks["independent_sql_number_present"] = all(value_present(result, row_value)
                for row in expected_rows for row_value in row if isinstance(row_value, (int, float)))
    if kind == "evidence_boundary":
        checks["boundary_answer_terms"] = all(term in answer for term in case.get("required_terms", []))
        checks["boundary_refusal_or_limit"] = any(term in answer for term in case.get("required_any_terms", [])) if case.get("required_any_terms") else True
        checks["no_unsupported_answer"] = not any(term in answer for term in case.get("forbidden_answer_terms", []))
    if kind.startswith("evidence") or kind.startswith("formula"):
        checks["required_source_documents"] = set(case.get("required_document_ids", [])).issubset(set(source_ids))
    if kind == "evidence":
        checks["grounded_terms_present"] = all(term in material for term in case.get("required_terms", []))
        checks["at_least_one_citation"] = bool(source_ids)
    if kind == "evidence_fact":
        checks["grounded_terms_present"] = all(term in material for term in case.get("required_terms", []))
    if kind == "evidence_boundary":
        checks["boundary_evidence_present"] = bool(source_ids)
    if kind == "formula":
        baseline = expected_rows[0][0]
        expected = round(float(baseline) * (1 + float(case["growth_rate"])), 2)
        checks["formula_value_present"] = value_present(result, expected)
        checks["formula_basis_present"] = value_present(result, baseline)
        checks["growth_rate_present"] = value_present(result, case["growth_rate"], tolerance=0.00001) or f"{case['growth_rate'] * 100:g}%" in material
    if kind == "formula_difference":
        baseline = expected_rows[0][0]
        expected = round(float(baseline) * (float(case["growth_rates"][0]) - float(case["growth_rates"][1])), 2)
        checks["difference_value_present"] = value_present(result, expected)
    if kind == "formula_pair":
        baseline = expected_rows[0][0]
        amounts = [round(float(baseline) * (1 + float(rate)), 2) for rate in case["growth_rates"]]
        checks["both_scenario_values_present"] = all(value_present(result, amount) for amount in amounts)
        checks["forecast_boundary_present"] = "实际销售额" in answer and any(term in answer for term in ("不能", "不可", "不是"))
    context_turns = data.get("context_turns", 0)
    checks["context_turns"] = context_turns
    checks["context_continuity"] = (context_turns == 0) if turn_index == 0 else (isinstance(context_turns, int) and context_turns > 0)
    checks["pass"] = all(value for key, value in checks.items() if key != "context_turns")
    return checks


def open_readonly(path):
    uri = path.resolve().as_uri() + "?mode=ro"
    connection = sqlite3.connect(uri, uri=True)
    connection.execute("PRAGMA query_only=ON")
    return connection


def main():
    suite = json.loads(QUESTIONS.read_text(encoding="utf-8"))
    output_dir = ROOT / "runtime"
    output_dir.mkdir(exist_ok=True)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    output = output_dir / f"authored-multiturn-25-{stamp}.json"
    with open_readonly(DATABASE) as db:
        workbook_growth = {}
        from openpyxl import load_workbook
        workbook = load_workbook(ROOT / "samples/documents/region-targets.xlsx", read_only=True, data_only=True)
        sheet = workbook.active
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            item = dict(zip(headers, row))
            if item.get("年份") == 2026:
                workbook_growth[item.get("地区")] = item.get("目标增长率")
        workbook.close()

        session_prefix = "auth25-" + uuid.uuid4().hex[:12]
        health_response = requests.get(BASE_URL + "/health", timeout=10)
        health_response.raise_for_status()
        health = health_response.json()
        docs_response = requests.get(BASE_URL + "/api/v1/knowledge/documents", timeout=30)
        docs_response.raise_for_status()
        loaded_docs = {item["document_id"] for item in docs_response.json().get("documents", [])}

        report = {
            "created_at": datetime.now(timezone.utc).isoformat(),
            "scope": suite["scope"],
            "suite_id": suite["id"],
            "base_url": BASE_URL,
            "health": health,
            "model": health.get("generation", {}).get("model"),
            "reasoning": "service configuration",
            "planned_turns": sum(len(group["turns"]) for group in suite["groups"]),
            "executed_turns": 0,
            "passed_turns": 0,
            "groups": [],
            "input_sha256": {str(QUESTIONS.relative_to(ROOT)).replace("\\", "/"): sha256(QUESTIONS),
                             str(DATABASE.relative_to(ROOT)).replace("\\", "/"): sha256(DATABASE),
                             **{str(path.relative_to(ROOT)).replace("\\", "/"): sha256(path) for path in SOURCE_FILES}},
            "loaded_document_ids": sorted(loaded_docs),
        }

        def save():
            output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

        for group in suite["groups"]:
            session_id = session_prefix + "-" + group["id"]
            group_report = {"id": group["id"], "title": group["title"], "session_id": session_id, "turns": []}
            report["groups"].append(group_report)
            for index, case in enumerate(group["turns"]):
                expected_rows = []
                if case.get("gold_sql"):
                    expected_rows = [list(row) for row in db.execute(case["gold_sql"]).fetchall()]
                elif case["kind"].startswith("formula"):
                    baseline_sql = ("SELECT SUM(sales_amount) FROM sales_orders WHERE region=? "
                                    "AND order_date>=? AND order_date<?")
                    year = int(case["year"])
                    expected_rows = [list(row) for row in db.execute(baseline_sql, (
                        case["region"], f"{year}-01-01", f"{year + 1}-01-01")).fetchall()]
                    if case["kind"] == "formula" and case["growth_source"] == "excel":
                        measured_rate = workbook_growth.get(case["region"])
                        if measured_rate is None or not math.isclose(float(measured_rate), float(case["growth_rate"]), abs_tol=1e-9):
                            raise ValueError("Excel source growth rate does not match frozen case: " + case["region"])
                started = time.perf_counter()
                record = {"turn": index + 1, "question": case["question"], "kind": case["kind"],
                          "expected_rows_independent": expected_rows}
                try:
                    response = requests.post(BASE_URL + "/api/v1/omni/query", json={
                        "question": case["question"], "session_id": session_id,
                        "reset_context": index == 0,
                    }, timeout=(10, 240))
                    record["http_status"] = response.status_code
                    data = response.json()
                    data["http_status"] = response.status_code
                    ids = document_ids(data)
                    checks = grade(case, data, expected_rows, ids, index)
                    record.update(response=data, citations_document_ids=ids, checks=checks, passed=checks["pass"])
                except Exception as exc:
                    record.update(passed=False, error_type=type(exc).__name__, error=str(exc)[:300])
                record["elapsed_seconds"] = round(time.perf_counter() - started, 3)
                group_report["turns"].append(record)
                report["executed_turns"] += 1
                report["passed_turns"] += int(record["passed"])
                save()
                summary = {"group": group["id"], "turn": index + 1, "passed": record["passed"],
                           "seconds": record["elapsed_seconds"],
                           "route": record.get("response", {}).get("route"),
                           "status": record.get("response", {}).get("status"),
                           "failed_checks": [name for name, value in record.get("checks", {}).items()
                                             if name not in {"pass", "context_turns"} and value is False]}
                print(json.dumps(summary, ensure_ascii=False), flush=True)
            group_report["passed_turns"] = sum(item["passed"] for item in group_report["turns"])
            group_report["total_turns"] = len(group_report["turns"])
            save()

        latencies = [turn["elapsed_seconds"] for group in report["groups"] for turn in group["turns"]]
        report["latency_seconds"] = {
            "mean": round(statistics.mean(latencies), 3) if latencies else None,
            "median": round(statistics.median(latencies), 3) if latencies else None,
            "p95": round(sorted(latencies)[math.ceil(0.95 * len(latencies)) - 1], 3) if latencies else None,
            "min": min(latencies) if latencies else None,
            "max": max(latencies) if latencies else None,
        }
        report["status"] = "passed" if report["passed_turns"] == report["planned_turns"] else "some_turns_failed"
        save()
    print(json.dumps({"status": report["status"], "passed_turns": report["passed_turns"],
                      "executed_turns": report["executed_turns"], "output": str(output),
                      "latency_seconds": report.get("latency_seconds")}, ensure_ascii=False), flush=True)
    return 0 if report["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
