"""Build anonymous contest materials from current source and evidence (no API call)."""
from __future__ import annotations

import hashlib
import importlib.metadata
import json
import platform
import re
import os
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'delivery'
FONT = 'Microsoft YaHei'


def evidence(name):
    return json.loads((ROOT / 'docs' / name).read_text(encoding='utf-8'))


def paragraph(doc, text, *, style=None):
    if style in {'Title', 'Subtitle', 'Heading 1', 'Heading 2'}:
        text = re.sub(r'[：:（）()、，·]', ' ', text)
        text = re.sub(r' +', ' ', text).strip()
    p = doc.add_paragraph(text, style)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.25
    return p


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Shading Accent 1'
    t.autofit = False
    if len(headers) == 4:
        widths = [3, 3, 5.5, 5.5]
    elif headers[0] in ('要求', 'ID'):
        widths = [1.5, 5.5, 10]
    else:
        widths = [3.8, 5.5, 7.7]
    for column, width in zip(t.columns, widths):
        column.width = Cm(width)
    properties = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement('w:' + edge)
        for key, value in {'val': 'single', 'sz': '4', 'color': 'D9D9D9'}.items():
            border.set(qn('w:' + key), value)
        borders.append(border)
    properties.append(borders)
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = str(text)
    for row in rows:
        for cell, text in zip(t.add_row().cells, row):
            cell.text = str(text)
    for index, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        for col, cell in enumerate(row.cells):
            cell.width = Cm(widths[col])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), '172A3A' if index == 0 else ('F3F6F8' if index % 2 else 'FFFFFF'))
            cell._tc.get_or_add_tcPr().append(shade)
            margins = OxmlElement('w:tcMar')
            for side in ('top', 'bottom', 'left', 'right'):
                margin = OxmlElement('w:' + side)
                margin.set(qn('w:w'), '100' if side in ('top', 'bottom') else '120')
                margin.set(qn('w:type'), 'dxa')
                margins.append(margin)
            cell._tc.get_or_add_tcPr().append(margins)
            for p in cell.paragraphs:
                if index == 0:
                    p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_after = Pt(5)
                p.paragraph_format.line_spacing = 1.1
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor.from_string('FFFFFF' if index == 0 else '000000')
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    return t


def save_report(stem, title, subtitle, pages):
    doc = Document()
    zoom = doc.settings.element.find(qn('w:zoom'))
    if zoom is not None:
        zoom.set(qn('w:percent'), '100')
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.core_properties.title = title
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21)
    sec.top_margin = sec.bottom_margin = Cm(1.85)
    sec.left_margin = sec.right_margin = Cm(2)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
        style = doc.styles[name]
        style.font.name = FONT
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
        style.font.color.rgb = RGBColor.from_string('000000')
        # The default Word Title style includes a blue rule. Remove it at
        # style level so WPS/Word do not reintroduce a decorative title border.
        for border in list(style._element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Heading 1'].font.size = Pt(19)
    doc.styles['Heading 2'].font.size = Pt(12)
    doc.styles['Title'].font.size = Pt(30)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('OmniAgent · ICT 赛题八  |  ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    paragraph(doc, '多模态数据驱动的可解释精准问数问答智能体', style='Subtitle')
    for _ in range(3):
        doc.add_paragraph()
    paragraph(doc, 'OmniAgent', style='Title')
    paragraph(doc, title, style='Heading 1')
    paragraph(doc, subtitle)
    paragraph(doc, '版本日期：2026-10-01\n匿名技术材料 · 依据官方赛题与当前实际运行证据')
    paragraph(doc, '阅读边界：本地开发验收与真实模型验收分别呈现。指定模型 gpt-6-luna 尚未通过鉴权；本文不把规则回退称为模型成功，不将开发题成绩换算为官方成绩。')
    for i, page in enumerate(pages):
        # Only the mandatory two highlight pages receive separate pages;
        # the body flows naturally instead of inflating length with white space.
        if i == 0 or (stem == '01-design-report' and i in (1, 2)):
            doc.add_page_break()
        paragraph(doc, page['title'], style='Heading 1')
        for item in page['items']:
            if isinstance(item, dict) and 'figure' in item:
                doc.add_picture(str(OUT/item['figure']),width=Cm(17))
                paragraph(doc,item['caption'])
            elif isinstance(item, tuple):
                table(doc, item[0], item[1])
            else:
                paragraph(doc, item)
    path = OUT / (stem + '.docx')
    doc.save(path)
    md = '# ' + title + '\n\n' + subtitle + '\n\n'
    for page in pages:
        md += '## ' + page['title'] + '\n\n'
        for item in page['items']:
            if isinstance(item, dict) and 'figure' in item:
                md += f"![{item['caption']}]({item['figure']})\n\n"
            elif isinstance(item, tuple):
                head, rows = item
                md += '| ' + ' | '.join(head) + ' |\n| ' + ' | '.join(['---'] * len(head)) + ' |\n'
                for row in rows:
                    md += '| ' + ' | '.join(str(x).replace('\n', '<br>') for x in row) + ' |\n'
                md += '\n'
            else:
                md += item + '\n\n'
    (OUT / (stem + '.md')).write_text(md, encoding='utf-8')
    return {'path': path.relative_to(ROOT).as_posix(), 'sections': len(pages), 'pagination':'office_export_required'}


def page(title, *items):
    return {'title': title, 'items': list(items)}


def main():
    OUT.mkdir(exist_ok=True)
    hybrid = evidence('HYBRID_ACCEPTANCE_REPORT.json')['summary']
    chinook = evidence('CHINOOK_RULES_REPORT.json')
    rag = evidence('RAG_SCALE_REPORT.json')
    probe = evidence('MODEL_API_PROBE.json')
    regression = evidence('LOCAL_REGRESSION_REPORT.json')
    faults = evidence('FAULT_RECOVERY_REPORT.json')
    first_faults = evidence('FAULT_RECOVERY_FIRST_RUN.json')
    if not regression['ok'] or faults['passed'] != faults['total']:
        raise ValueError('当前回归或故障审计未通过，不能更新材料')
    test_count = regression['passed']
    large = next(r for r in rag['records'] if r['pages'] == 500)['global_semantic_query']
    perf_text = (f"500页全局语义热P50={large['warm']['p50_ms']}ms、P95={large['warm']['p95_ms']}ms，"
                 f"首次全局回答{large['first_answer_ms']/1000:.3f}s。")
    fault_text = (f"同一组十项实际临时副本故障首次{first_faults['passed']}/{first_faults['total']}，"
                  f"修复后{faults['passed']}/{faults['total']}；原文件篡改或缺失停止旧切片与公式链，"
                  'SQL缺库/锁定停止后步，硬kill子服务重启后五轮和pending澄清继续，'
                  '会话锁503后解锁保留原历史。仅为开发审计，不代表掉电容灾。')
    flow_rows = [('客单价','文档公式→SQL销售额/订单数→计算','来源与零分母'),
                 ('预测','PDF公式→SQL基准→Excel增长率','年份与单位'),
                 ('地区问数','Excel地区→SQL过滤','实际单元格与地区'),
                 ('冠军经验','SQL排名→冠军实体→文档检索','实体与引用'),
                 ('阈值比较','检索→来源事实→Excel→le比较','原文、数值与单位')]
    qa = f"{hybrid['qa_pass']}/{hybrid['qa_total']}"
    sql = f"{hybrid['sql_pass']}/{hybrid['sql_total']}"
    fusion = f"{hybrid['fusion_pass']}/{hybrid['fusion_total']}"
    perf = [(str(r['pages']), str(r['chunks']), str(r['global_semantic_query']['warm']['p50_ms']),
             str(r['global_semantic_query']['warm']['p95_ms'])) for r in rag['records']]
    modules = [
        ('结构化问数', 'backend/nl2sql/', 'Schema与值链接→计划→安全SQL→只读执行'),
        ('文档知识库', 'knowledge_store.py / chunk_cleaning.py', '原文件SHA→解析/OCR→有定位的切片'),
        ('混合检索', 'dense_retrieval.py / cross_source.py', 'BM25与本地BGE→RRF→编号/证据门'),
        ('跨源推理', 'dependency_agent.py / formula_binding.py', '有界依赖计划→来源数值→公式与单位核验'),
        ('可信生成', 'grounded_generation.py / responses_client.py', '结构化claims→原文引用与数字核对'),
        ('持续会话', 'session.py / omni_agent.py', '结构化槽位→五轮上下文→主题清空'),
    ]
    intermediate = [
        ('M01', '要素识别、别名与改写', 'NFKC、值词边界、5/5文本扰动'),
        ('M02', '低示例跨领域', 'Chinook 11表、0 few-shot，开发题12/12'),
        ('M03', '主动澄清', '指标、时间、值、JOIN歧义与回填'),
        ('M04', '数十表Schema Linking', '11/40/80表干扰，9/9开发案例'),
        ('M05', 'JOIN、聚合、嵌套', '关系图、复合键、粒度与防扇出'),
        ('M06', '文档公式与数据结合', 'AST、来源绑定、单位与预测年份守卫'),
        ('M07', '非标准目录', '8份实际PDF首次2/8→8/8；原文件行号'),
        ('M08', '复杂度自适应', 'PageSignal及900字/80重叠复杂切片'),
        ('M09', '文档质量综合评估', 'OCR8/10→10/10；繁体/错字待验收'),
    ]
    claims = [
        ('问数开发验收', sql, '规则规划、合成开发题'),
        ('多格式问答开发验收', qa, '混合检索、原文摘录与引用核对'),
        ('跨源工具链开发验收', fusion, '明确工具计划，非未知模型规划'),
        ('公开Chinook开发题', f"{chinook['passed']}/{chinook['total']}", '公开库，人工题，非Spider基准'),
        ('实际故障开发审计', f"{first_faults['passed']}/{first_faults['total']} → {faults['passed']}/{faults['total']}", '临时文件/数据库/HTTP子服务，首次失败保留'),
        ('真实gpt-6-luna', probe['status'], 'models及responses探测401，未有模型成绩'),
    ]
    design = [
        page('核心亮点预览（一）：架构与证据',
             '系统面向“数据库里的数值”和“文档里的口径”需要共同回答的问题。结构化问数、独立知识库与跨源执行器共享来源契约；前端直接展示实际SQL、引用位置、计算参数和失败原因。',
             ({'figure':'architecture.png','caption':'图1 系统架构与共享证据契约（与当前源码对应）'}
              if (OUT/'architecture.png').exists() else (['层次', '实现位置', '数据流'], modules)),
             '处理链路：用户问题→结构化会话→SQL/文档/跨源路由→有界工具执行→证据与单位校验→答案和可回看来源。PDF、DOCX、XLSX、TXT、MD、PNG入库后保留页、行、单元格或OCR区域定位。',
             (['当前验收', '结果', '口径'], claims)),
        page('核心亮点预览（二）：中级任务覆盖',
             (['要求', '能力', '当前证据'], intermediate),
             '重点贡献是把公式、数值、单位、日期版本及工具依赖绑定到真实来源。缺失依据、政策重叠、零分母、币种冲突和已知预测年份不符均停止最终结论。',
             '覆盖实现不代表所有官方未知测试均已通过。模型通用规划、独立保留题、公开OCR基准及跨源五轮真实模型评测仍需进一步验证。'),
        page('需求与任务边界',
             '依据 specification/official-track8.pdf（14页正式赛题），基础任务包括单表问数、单文档问答、可解释展示及交互Demo；中级任务涵盖九类难点；决赛需要多源协同、多跳、澄清与连续五轮。',
             '本项目建立在独立目录中，迁移已有NL2SQL与未提交优化作为受保护基线。原生产项目仅作为只读借鉴源，不作为运行服务、私有知识库或密钥依赖。',
             (['对象', '当前支持', '边界'], [('数据库','SQLite演示库与公开Chinook','其他数据库方言未作运行验收'),('文档','15份实际文件、6种格式','合成样本，非真实客户资料'),('模型','指定Responses兼容接口','仅gpt-6-luna，鉴权尚未通过'),('交互','问数/知识库/验收工作台','本地开发服务，不声称完成生产审计')]),
             '系统对无证据、歧义、错误计划和不可执行的查询显式返回状态；不会为提高表面完成率而使用猜测答案。'),
        page('问数设计：意图与可执行口径',
             'Schema自动解析表、字段、主外键与关系；值索引把“华东”等实际值绑定到正确列。Unicode归一化、繁简与别名处理降低输入差异；显式字段与最长字段匹配防止短词抢占长指标。',
             '计划由指标、时间、过滤、分组、聚合、排序与JOIN结构组成。模型不直接执行裸SQL；同一安全门检查字段合法性、覆盖率、JOIN粒度和只读性。模糊意图进入澄清接口，用户选择后回填。',
             '跨域扩展首先增加Schema和业务指标口径，而非为每一道题追加字符串模板。Chinook与干扰表实验证明有限场景可运行，尚不能推导任意业务数据库准确率。',
             '五轮会话示例：2025华东销售额→改华南→改2024→换订单数→换华北。结构化条件替换避免把前轮已撤销条件误叠加。'),
        page('多模态与检索设计',
             '文本PDF优先保留原生文本与页块，扫描PDF和图片调用RapidOCR ONNX CPU。DOCX段落、XLSX行列、TXT/MD标题路径保留结构。低质量图片执行有界增强并保留原始结果。',
             'BGE-small-zh-v1.5固定公开revision与SHA，以CLS归一化向量提供真正Dense检索。BM25提供精确词匹配，RRF融合排名。明确编号先作为范围约束，避免CASE0005与CASE00050混淆。',
             '知识库原文件内容寻址，保存SHA与定位片段。答案可以回看原文件；定位、原文和数字核验失败时拒绝把候选片段包装为确定结论。',
             '开发阈值0.35/0.6未经过独立校准；OCR置信度选优是工程启发式，不是语义正确保证。'),
        page('跨源设计：每一步都有来源',
             '执行计划每步包含id、tool和args，最多16步；ref/path定义实际依赖。运行前拒绝循环、缺失依赖、重复ID和未知工具。SQL或文档结果只有通过契约后才能作为后步参数。',
             (['场景', '依赖链', '核验'], flow_rows),
             '已核对例子：2025华东销售额29584，按来源增长率12%得到2026预测33134.08。错年份停止，不显示该中间数值为最终答案。该结果仅用于合成样本验收。'),
        page('评估设计与当前结果',
             (['评测', '结果', '解释'], claims),
             f'本地{test_count}项回归覆盖接口、安全、语义与拒绝路径；一个Starlette弃用警告不影响当前通过，但需在依赖升级时处理。',
             fault_text,
             '结果匹配优先比较执行后的行与列，不把SQL字符串相等当作准确率。问答同时检查预期事实与引用；OCR对比原图与同一扰动样本，避免无对照的增强宣传。',
             '计划下一阶段冻结独立保留题、引入未知Schema、遮蔽开发答案后实测；模型验证必须检查provider/model及回退状态，不允许混入规则结果。'),
        page('效率、成本与工程取舍',
             (['PDF页数','切片数','全局语义热P50/ms','全局语义热P95/ms'],perf),
             '以上是实际文字PDF入库后的混合检索与摘录答案，不含OCR与外部生成。' + perf_text + '首次全局编码、首载模型与热缓存分别报告；当前首次全局并不是新进程完整冷启动。编号查询不代表一般语义延迟。',
             '数据库已实测1千、1万、10万行；10万行热简单查询约11.6ms、嵌套约77.2ms、同比约291.7ms。1/4/8线程规则查询约13.63/28.44/33.93QPS，不当作模型吞吐。',
             '优化优先顺序：缓存Schema/向量、预热模型、索引和限定候选，再控制工具调用与上下文长度。API计费和端到端token成本因401尚未实测，不给虚构费用。'),
        page('交付与部署',
             '完整资产包包含源码、Schema、15份实际样本、公开BGE权重、Chinook与许可。runtime凭据、日志、缓存及基线ZIP不进入交付包。依赖锁定文件记录当前实测环境，跨平台安装仍需对应平台wheel。',
             '启动：python -m pip install -r delivery/requirements-tested.txt；python tools/run_server.py --port 8030。打开首页、knowledge.html和capabilities.html。API模式先probe通过，再加--with-model；未通过时保留显式规则/摘录模式。',
             '验包工具核对逐文件SHA，在全新目录解包并重建数据库、样本入库、随机端口启动，验证health、前端、SQL、Dense、RAG与OCR。当前检查点通过，最终材料加入后需再次验包。',
             '组委会Word/PPT模板尚未提供，本材料采用匿名自定义排版，不声称已套用官方模板。最终提交前应按届时模板迁移，并补录真实模型及保留集结果。'),
    ]
    technical = [
        page('系统架构与运行模式', (['模块','源码','契约'],modules),
             'FastAPI提供统一同源接口，开发服务绑定127.0.0.1；前端不读取密钥。用户通过omni/query进入结构化SQL、文档或融合计划。规则模式、模型模式和回退状态在返回结果中区分。'),
        page('Schema Linking与要素识别',
             'backend/nl2sql读取SQLite元数据，构建表字段与主外键图；值词与字段别名共同产生候选。列名采用NFKC归一化，保留原始SQL标识符并统一引用，避免中文/英文或全角输入导致不可执行。',
             '显式字段优先，长字段优先于短字段；主题日期与业务时间槽位分别处理。对于“金额”“数量”等宽泛指标先展示候选而非擅选。值边界避免词中片段命中错误实体。',
             '11/40/80表的9个开发测试核对选择的表、字段、结果，Schema干扰并非真实企业80表基准。角色关联与同名列的泛化仍需要额外盲测。'),
        page('查询计划、JOIN与SQL验证',
             '自然语言解析得到可校验结构化计划：目标指标、聚合、分组、时间窗口、过滤与关系路径。模型仅提交受限结构，编译器处理参数化值，禁止模型生成SQL绕过编译。',
             '关系图检查完整复合键和连接方向。聚合前确认事实表粒度，拒绝会让数值复制的JOIN；多种同等关系时请求澄清。去重计数与合计保持不同语义。',
             'SQL AST和只读连接拒绝DDL/DML、多语句、未批准字段及危险结构；限制行数和执行时间。空结果作为有效结果状态，区别于数据库错误或缺少证据。',
             '用金标SQL执行结果核对，而不是词面模板吻合。Chinook12道题覆盖已知场景，复杂任意SQL生成能力尚未取得独立评测成绩。'),
        page('澄清与持久会话',
             'nl2sql/clarify接收明确候选选择并回填计划。澄清针对会改变结果的字段、单位、时间或连接歧义；安全错误不得通过用户随意点击绕过。',
             'session.py以SQLite保存结构化槽位与最近五轮状态，设置TTL与容量；重启可以恢复。新主题清空不适用条件，短跟进输入替换条件，不把全部历史自然语言简单拼接。',
             '网页已经完成五轮地区/年份/指标替换，结果29584→22992→4998→1→1。实际kill子服务并重启后相同序列和pending澄清继续；会话库锁返回脱敏503和Retry-After: 2，不创建新内存身份。该序列是合成规则模式验收，不代表多源模型五轮未知题已经全部通过。'),
        page('文件解析、结构与切片',
             'knowledge_store.py保存文档元数据、内容地址与切片。支持PDF、DOCX、XLSX、TXT、MD、PNG；先检查格式、大小与解析并发，再进入格式专用解析器。',
             'text_structure.py识别Markdown、数字、中文与括号标题，保存行号与父级路径。标题跳跃给出告警，不把规则推断层级标作人工金标。',
             'PageSignal记录页上文本、块、图及质量等实际信号，决定解析与切片策略。复杂路径采用900字长度和80字重叠；正文、表格与OCR区域都有对应locator。',
             '原文件下载、document读取、公式与最终检索命中的来源均重新核验SHA及路径；不缓存校验通过，旧切片和向量不能使篡改/缺失原文件继续充当证据。仅校验选中来源，减少全库哈希。知识库和原始样本属于独立项目。'),
        page('OCR与低质量恢复',
             'ocr.py使用RapidOCR ONNX CPU识别中文与英文，返回文字、bbox、置信度及变换历史。扫描PDF按页栅格化，图片保留来源。',
             '质量判断同时参考图像信号与识别结果。低质量图片即使原始平均confidence偏高也执行有界增强；最多三次尝试，将原图也保留为候选，选择置信度较高的结果。',
             '十种成对合成扰动原图8/10，pipeline10/10。固定预期“2小时/24小时”等事实核验，不只比较置信度。该实验非OHR-Bench，样本量不足以给统计泛化保证。',
             '风险：高置信错字、表格错行及编号误读仍可能出现。对于影响计算的OCR数值保留原始区域与低置信提示，未知金额不可无来源写入工具参数。'),
        page('Dense、BM25与融合检索',
             'dense_retrieval.py加载BGE-small-zh-v1.5，revision=7999e1d3359715c523056ef9478215996d62a620，使用safetensors、不允许remote code或pickle加载。CLS向量归一化后执行cosine计算。',
             'BM25与Dense结果按RRF融合并保留各路rank/cosine/score。缓存key包含模型revision、标题和内容；更新文档或模型不能误复用旧向量。',
             '明确实体编号作为范围约束：CASE0005必须精确出现，CASE00050不满足；不存在的编号返回无证据。这是精确实体守卫，不宣称通用实体链接已解决。',
             '开发阈值属于待校准参数。候选片段不足时摘录回退或拒答，不能把高相似度当作事实蕴含。'),
        page('有依据生成与指定API',
             'responses_client.py提供Responses兼容结构化输出、超时和脱敏audit。grounded_generation.py要求claims与literal quote，核验引用存在及数字一致。逐字引用有效并不证明完整语义蕴含。',
             'tools/model_runtime.py仅读取runtime/model_config.json；模型只允许gpt-6-luna，域名只允许https://spacetimeai.cc（实际请求/v1）。不读Codex旧凭据，不切换其他模型。',
             '指定API实际探测：/v1/models、/models和/v1/responses均401。models失败不能直接判断模型不存在；401只说明请求未通过鉴权。当前没有可用的真实生成成绩。',
             '鉴权恢复后依次运行probe_model.py与evaluate_model.py --full，检查返回model、provider、usage及回退标签。禁止用成功的规则SQL替代真实模型能力证明。'),
        page('有界DAG与跨源工具',
             'dependency_agent.py支持sql、search、search_fact、document_formula、cell、fact、calculate、policy_select、compare。每一步ID唯一；ref/path形成实际依赖图，上游失败时下游不执行。',
             '计划最大16步；先验证依赖、循环、工具白名单与参数契约。运行trace记录输入证据、结果与状态，前端展示实际调用链，不展示虚构的模型内部思考。',
             '工具输出分离数值、单位、来源与验证状态。数据型参数只能引用SQL单元格、XLSX或已定位文档，不接受literal代替可信来源。',
             '融合五流程已通过明确计划验收；未知自然语言转复杂计划的鲁棒性需gpt-6-luna实测。'),
        page('实际故障注入与恢复',
             fault_text,
             '证据完整性、存储不可用与工具契约错误分别返回evidence_integrity_failed、storage_unavailable、tool_contract_failed。failed_task、skipped_tasks和真实edges保留，前端不显示最终成功答案。',
             'SQL engine使用contextmanager加finally close，初始化连接使用closing；事务with结束不能代替连接关闭。正常和异常退出均有立即关闭回归。',
             '首次4/10与当前10/10保存相同用例签名；只注入临时合成副本，不改主服务或真实数据。13项新增故障回归包括损坏数据库、坏PDF替换、恢复后查询与数据变更。'),
        page('公式、单位、年份与政策版本',
             'formula_binding.py从实际文档提取公式并解析AST，白名单算术拒绝代码执行与非法表达式。零分母和非有限数值立即失败，参数名称与引用来源显式绑定。',
             'unit_algebra.py管理数值尺度、百分数、金额及币种，禁止冲突币种默默计算。年份守卫核验文档基准年/目标年、SQL全年范围以及Excel参数适用年；无法推断时标not_inferred。',
             'policy_evidence.py按明确生效日期选版本，重叠或缺失阻止答案。已验证退货期限5日→7日的边界，ASCII逗号不再让政策值带入整个日期句。',
             '未知单位、隐含生效条件和复杂公式语义仍可能需要主动澄清，不由启发式强行补齐。'),
        page('解释链路与前端契约',
             '首页展示问数、SQL与结果；knowledge.html展示文档、统一多轮及跨源场景；capabilities.html按赛题要求列出证据与未完成项。',
             '可解释信息包括SQL及参数、来源SHA、页/行/单元格/OCR定位、公式、引用值、单位与日期、工具依赖和失败原因。提供的解释是执行记录，不承诺揭示模型隐含推理。',
             '融合结果incomplete时不把已完成步骤的中间数值当作最终答案；模型生成失败保持显式摘录回退。API探测状态显示时间戳，保存的历史记录不等于实时可用性。',
             '服务限制查询长度、行数、文件大小与并发；生产Bearer/CORS配置有边界。当前本地部署尚不等同多租户生产权限审计。'),
        page('功能评估：输入、金标与判定',
             (['测试','当前结果','判定口径'],claims),
             f'本地{test_count}项回归主要检验实现契约、异常与安全路径；不能把单元测试数量作为准确率分母。开发问数比较结果，问答核对事实与原文，跨源核对数值及来源依赖。差旅新Schema首次6/8修复后8/8，PDF目录首次2/8修复后8/8，首次失败保留。',
             '评测输入、结果、运行脚本和范围标记一起入包。history中的旧模型结果只属于历史资料，不纳入gpt-6-luna评测。',
             '未见题验收需要冻结输入/金标且避免修复后仍称其为未见；报告失败原因为Schema、检索、生成、计算或授权，不能只输出平均成功率。'),
        page('复杂度与规模效率',
             (['页数','切片','语义热P50/ms','语义热P95/ms'],perf),
             '设表数T、字段数C、关系边E、切片数N、向量维度d、检索k、计划步数S。Schema遍历约O(T+C+E)，关系路径搜索约O(T+E)；SQL成本由扫描、索引与JOIN基数决定。',
             '当前精确Dense矩阵扫描约O(Nd)，Top-k约O(N log k)；BM25倒排依赖命中posting数。RRF融合约O(k log k)，DAG校验O(S+依赖边数)。外部模型延迟随输入输出token变化，不能套用本地线性公式。',
             '1k/10k/100k行与10/100/500页已经分规模实测。冷编码、热缓存、编号缩小候选与全局语义必须分别报告；OCR和真实模型端到端效率尚未验收。'),
        page('复现、许可与未完成项',
             'delivery/requirements-tested.txt记录直接依赖实测版本；ENVIRONMENT.json记录完整包版本、Python与平台。模型权重与Chinook各有ASSET_MANIFEST和LICENSE，样本为自建合成CC0。',
             '先安装依赖，按README启动；完整资产包无需复制原项目。verify_checkpoint.py对新目录解包、逐文件哈希、重建与独立随机端口作验收，源码包则需要公开资产下载。',
             '未完成：gpt-6-luna有效鉴权；独立未知题；公开复杂OCR样本；通用跨源五轮；真实模型耗时/token成本；官方模板迁移及真实评审反馈。',
             '本文按官方赛题内容制作，不以未提供的官方模板名义提交；仅有当前证据的条目可表述为“已验收”。'),
    ]
    deep = [
        page('决赛深化范围与证据等级',
             '本材料是当前实现的决赛技术深化准备稿，不假称已经收到初赛评审反馈。最终版应补充真实评委意见、逐项响应与新增实验。',
             (['等级','定义','当前情况'],[('实现','源码存在且通过契约回归','九类中级能力与融合执行器'),('开发验收','已知输入与金标实测','SQL、QA、融合、Chinook与扰动'),('独立效果','保留题或公开标准协议','尚未提供完整成绩'),('真实模型','指定接口成功且无回退','gpt-6-luna鉴权401')]),
             '评委可以在验收工作台逐项回看证据，所有指标附范围；不把覆盖数量、开发题通过率和奖项预期混在一起。'),
        page('多源任务组合与现场路线',
             '五种已通过的跨源流程包含文档公式+SQL客单价、PDF+SQL+Excel预测、Excel地区→SQL、SQL排名→文档冠军经验和文档检索+Excel阈值。政策版本比较是额外边界案例。trace核对前步结果对后步的实际影响。',
             '主路线：自然语言问数→文档解释→三源预测→故意错年份→政策新旧日期→澄清选择→五轮条件变化。失败案例和正常案例同屏展示。',
             '当前工具计划可运行；通用模型计划仍需指定接口恢复后验证。现场无API时明确进入规则/摘录演示，不使用预录模型答案冒充在线结果。'),
        page('可计算证据与时间一致性',
             '三源预测必须具有文档公式、SQL基准数值与参数表适用年。每项带locator和SHA，计算器保留AST与引用。',
             '29584×(1+12%)=33134.08为合成验收例子；若目标年/基准年/SQL全年窗口不匹配，返回incomplete。未知年份标not_inferred，不默认为验证通过。',
             '政策版本按问题日期选择已生效条款；重叠条款不能靠最新文档名猜测。退货期限5日和7日已作边界比较，依赖日期证据而不是文件排序。'),
        page('可靠性深化：拒绝也是正确行为',
             (['故障','期望行为','证据'],[('未知实体编号','无证据，不用近似编号','编号边界回归'),('参数缺失/零分母','不输出最终数值','DAG与AST测试'),('JOIN粒度不符','澄清或拒绝','复合键/防扇出测试'),('币种冲突','停止计算','单位契约'),('API鉴权失败','明确回退或停止','401脱敏探测')]),
             fault_text,
             '正确拒绝率和错误拒绝率应在后续独立题集分别报告。仅测正常路径会掩盖系统在数据冲突时生成自信错误的风险。'),
        page('检索深化：语义召回与实体边界',
             'Dense可以弥补纯词匹配，但近似编号和不同政策版本必须额外控制。实体约束与时间约束先限定可比证据，RRF再在合法范围中融合。',
             '第一次500页评测暴露了错误合同编号，已通过精确编号守卫修复。这项修复说明需要错误案例反馈，不意味着Dense所有语义检索问题已经解决。',
             '编号和无编号全局语义评测分开。' + perf_text + '来源选中后核验；首次全局、首载模型与热缓存单列，不以编号快速路径替代一般语义延迟。新旧环境/缓存不同，不宣称固定倍数加速。'),
        page('低质量文档的证据保留',
             'OCR返回识别原文、区域、confidence与变换历史，允许回看原图。三次尝试的选择规则保留原始候选，避免增强覆盖正确识别。',
             '合成成对扰动8/10→10/10仅为开发证据。下一阶段需要真实拍照、倾斜、污损、表格错行与不同扫描仪，按文档分组拆分，防止同一页的不同扰动跨越训练与测试。',
             '评价应包含字符错误率、关键数字准确率、引用定位与下游问答结果。平均置信度只用于候选排序，不作为最终准确率。'),
        page('未知题与消融实验方案',
             '冻结保留题：问数按未见Schema、表达和SQL结构划分；文档按未见原文件划分；融合按未见依赖组合划分。金标在实验前固定，并标明制定者与来源。',
             (['消融','比较对象','指标'],[('移除Dense','BM25与混合检索','Recall@k、来源准确率'),('移除编号守卫','近似编号冲突题','误引用率'),('移除年份/单位守卫','冲突融合题','危险计算率'),('关闭OCR增强','同文档成对扰动','关键事实正确率'),('关闭结构会话','五轮条件变化','每轮执行准确率')]),
             '这些是待执行的实验方案，不预填增益或显著性。无需宣称优于公开SOTA；先证明对目标场景的可复现效益。'),
        page('部署、成本与生产前置',
             '当前包含本地95,827,648字节BGE权重与Chinook，在线生成依赖指定API。CPU检索可独立运行；OCR与模型首载占用单独统计。',
             '生产前需要数据权限、租户隔离、审计、配置轮换、限流与故障恢复。现有只读SQL和开发本机绑定解决执行边界，不等同数据库所有表都可以向所有用户开放。',
             '真实评测准备44题，SQL内部也接真实规划器；每题汇总路由、生成、SQL及重试，usage缺失保持未知。唯一预检401，题目未执行，不估费用。压缩包不是无依赖离线安装镜像。'),
        page('材料依据与最终提交门槛',
             '全部量化结论对应docs中的原始JSON，delivery/SOURCE_MANIFEST.json保存本轮材料来源SHA。报告和PPT不包含用户名、单位、密钥或真实客户信息。',
             '最终提交门槛：有效gpt-6-luna实测；五类跨源与五轮模型验收；独立保留集；依赖锁；材料页数与匿名检查；完整资产新目录启动；逐文件哈希；官方模板检查。',
             '当前已经实现的能力可以演示，但这些未完成门槛必须如实列明。正式定稿更新资料与实验数据后重新生成材料和压缩包，不能沿用过期报告。'),
    ]
    innovation = [
        page('定位与核心贡献',
             '申报方向：实用创新与转化价值。当前没有公开最优结果的充分对比证据，不申报已经证明的技术领先性。',
             '核心贡献是可计算证据契约：公式来自文档，参数来自SQL或单元格，单位与年份必须一致，政策按生效日期选择。每项输入可回看到原始来源，冲突时停止而不拼凑答案。',
             '目标场景为企业问数与文档口径共同决定答案的流程。单独的SQL或检索都不足以可靠回答“按新政策、用当年参数预测”的问题。',
             '有界依赖执行器将这些要求落实为可检查的工具接口，最多16步，并为故障、歧义与无依据结果提供明确状态。'),
        page('举证材料与成对对照',
             (['证据','结果','范围'],claims[:4]),
             'OCR原图8/10与增强10/10来自同一组十种合成扰动；编号守卫来自500页检索中实际暴露的错编号；年份守卫来自跨源预测边界。',
             fault_text,
             '上述证据证明实现可以复现并具备可解释失败路径，不能证明业界领先。独立保留题、公开复杂OCR和模型规划仍待验证。'),
        page('场景适配、效率与价值边界',
             '系统可在本地CPU进行公开BGE检索与真实OCR，默认规则/摘录允许无API演示。内容寻址、版本清单和精确依赖降低交付时“只在开发电脑可跑”的风险。',
             perf_text + '数据规模、首次全局编码和热缓存分开，首次全局并不是新进程完整冷启动；不含OCR。模型生成耗时及费用尚未实测。',
             '潜在价值是降低数值与口径错配、提升复查效率；当前没有真实用户节省工时或营收改善证据，不给虚构商业收益。',
             '评审价值在于来源、计算与失败边界可验证，而非仅展示流畅回答。未来应通过真实授权数据、独立盲测与用户任务时长证明转化收益。'),
    ]
    outputs = [
        save_report('01-design-report','设计报告','设计目标、核心亮点、中级任务与验收边界',design),
        save_report('02-technical-specification','技术实现说明书','模块契约、实现机制、评测与复现',technical),
        save_report('03-finals-technical-report','决赛深化技术文档（准备稿）','多源推理、可靠性、实验设计与提交门槛',deep),
        save_report('04-innovation-statement','突破与创新性自述','实用创新与转化价值 · 当前可核验贡献',innovation),
    ]
    # Author DOCX with the bundled runtime while recording the interpreter
    # actually used for the app/tests, rather than substituting authoring deps.
    tested_python = os.environ.get('ICT8_TESTED_PYTHON', sys.executable)
    if Path(tested_python).resolve() != Path(sys.executable).resolve():
        snapshot = subprocess.run([tested_python, '-c',
            'import importlib.metadata,json,platform; print(json.dumps({"python":platform.python_version(),"platform":platform.platform(),"machine":platform.machine(),"packages":dict(sorted((d.metadata["Name"],d.version) for d in importlib.metadata.distributions() if d.metadata.get("Name")))}))'],
            capture_output=True, text=True, check=True)
        tested_env = json.loads(snapshot.stdout)
    else:
        tested_env = {'python':platform.python_version(),'platform':platform.platform(),'machine':platform.machine(),
                      'packages':dict(sorted((d.metadata['Name'],d.version) for d in importlib.metadata.distributions() if d.metadata.get('Name')))}
    installed = {name.lower().replace('_','-'):version for name,version in tested_env['packages'].items()}
    requirements = []
    uninstalled = []
    for line in (ROOT/'ict-track8/requirements.txt').read_text(encoding='utf-8').splitlines():
        name = re.split(r'[<>=!~\[]', line.strip())[0]
        if name:
            version = installed.get(name.lower().replace('_','-'))
            if version is None:
                uninstalled.append(name)
                continue
            extras = '[standard]' if name == 'uvicorn' else ''
            requirements.append(f'{name}{extras}=={version}')
    (OUT/'requirements-tested.txt').write_text('\n'.join(requirements)+'\n',encoding='utf-8')
    env = {'created_at':datetime.now(timezone.utc).isoformat(),**tested_env,
           'scope':'observed_environment_not_cross_platform_resolver_lock',
           'requirements_not_installed':uninstalled}
    (OUT/'ENVIRONMENT.json').write_text(json.dumps(env,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    sources = []
    for source in sorted((ROOT/'docs').glob('*REPORT.json')) + [ROOT/'docs/FAULT_RECOVERY_FIRST_RUN.json',ROOT/'docs/DOMAIN_TRANSFER_FIRST_RUN.json',ROOT/'docs/PDF_OUTLINE_FIRST_RUN.json',ROOT/'docs/FAULT_RECOVERY_UPDATE_20261001.md',ROOT/'docs/MODEL_API_PROBE.json',ROOT/'specification/official-track8.pdf',ROOT/'docs/CURRENT_ARCHITECTURE.md',ROOT/'docs/ACCEPTANCE.md']:
        if not source.exists():
            raise ValueError('材料依据文件不存在: ' + source.name)
        sources.append({'path':source.relative_to(ROOT).as_posix(),'sha256':hashlib.sha256(source.read_bytes()).hexdigest()})
    (OUT/'SOURCE_MANIFEST.json').write_text(json.dumps({'generated_at':datetime.now(timezone.utc).isoformat(),'sources':sources,'documents':outputs},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'outputs':outputs,'environment':'delivery/ENVIRONMENT.json'},ensure_ascii=False))


if __name__ == '__main__':
    main()
