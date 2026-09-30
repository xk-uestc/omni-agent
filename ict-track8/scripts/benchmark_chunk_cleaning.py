"""Benchmark deterministic PDF, DOCX, and XLSX chunk parsing in memory."""

from __future__ import annotations

import argparse
import json
import math
import statistics
import sys
import time
import tracemalloc
from io import BytesIO
from pathlib import Path
from typing import Callable, TypeVar

import fitz
from docx import Document
from openpyxl import Workbook

PACKAGE_ROOT = Path(__file__).resolve().parents[1]
if str(PACKAGE_ROOT) not in sys.path:
    sys.path.insert(0, str(PACKAGE_ROOT))

try:
    from backend.chunk_cleaning import DocumentChunker
except ModuleNotFoundError:
    sys.path.insert(0, str(PACKAGE_ROOT))
    from backend.chunk_cleaning import DocumentChunker

ResultT = TypeVar("ResultT")


def _pdf_bytes(pages: int) -> bytes:
    document = fitz.open()
    for index in range(pages):
        page = document.new_page(width=612, height=792)
        page.insert_text((72, 72), f"Synthetic page {index + 1}: searchable operational evidence.")
        page.insert_text((72, 110), f"Record count {index + 1} and quality threshold 0.95.")
    return document.tobytes()


def _docx_bytes(paragraphs: int) -> bytes:
    document = Document()
    for index in range(paragraphs):
        document.add_paragraph(f"Synthetic paragraph {index}: operational evidence and traceable source data.")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _xlsx_bytes(rows: int) -> bytes:
    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet("Synthetic")
    worksheet.append(("record_id", "category", "amount", "active"))
    for index in range(rows):
        worksheet.append((index + 1, f"category-{index % 20}", index * 1.25, index % 2 == 0))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _measure(parser: Callable[[], ResultT], modality: str, scale: int, iterations: int) -> dict[str, object]:
    samples = []
    peak_bytes = 0
    chunk_count = 0
    for _ in range(iterations):
        tracemalloc.start()
        started = time.perf_counter()
        result = parser()
        samples.append((time.perf_counter() - started) * 1000)
        _current, peak = tracemalloc.get_traced_memory()
        tracemalloc.stop()
        peak_bytes = max(peak_bytes, peak)
        chunk_count = len(result.chunks)
    ordered = sorted(samples)
    p95_index = max(0, min(len(ordered) - 1, math.ceil(len(ordered) * 0.95) - 1))
    return {
        "modality": modality,
        "input_scale": scale,
        "iterations": iterations,
        "chunk_count": chunk_count,
        "p50_ms": round(statistics.median(ordered), 3),
        "p95_ms": round(ordered[p95_index], 3),
        "peak_tracemalloc_bytes": peak_bytes,
    }


def benchmark(scale: int, iterations: int) -> list[dict[str, object]]:
    pdf = _pdf_bytes(scale)
    docx = _docx_bytes(scale)
    xlsx = _xlsx_bytes(scale)
    chunker = DocumentChunker()
    return [
        _measure(lambda: chunker.parse_pdf(pdf, document_id=f"bench-{scale}.pdf"), "pdf_pages", scale, iterations),
        _measure(lambda: chunker.parse_docx(docx, document_id=f"bench-{scale}.docx"), "docx_paragraphs", scale, iterations),
        _measure(lambda: chunker.parse_xlsx(xlsx, document_id=f"bench-{scale}.xlsx"), "xlsx_data_rows", scale, iterations),
    ]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--scales", nargs="+", type=int, default=[10, 100, 500])
    parser.add_argument("--iterations", type=int, default=3)
    parser.add_argument("--json-out", type=Path)
    args = parser.parse_args()
    if args.iterations < 1 or args.iterations > 20:
        parser.error("--iterations must be between 1 and 20")
    if any(scale < 1 or scale > 5000 for scale in args.scales):
        parser.error("--scales must be between 1 and 5000")
    report = {
        "benchmark": "ict-track8-chunk-cleaning",
        "scope": "synthetic_in_memory_parsing_not_official_dataset_or_production_upload_path",
        "results": [item for scale in args.scales for item in benchmark(scale, args.iterations)],
    }
    rendered = json.dumps(report, ensure_ascii=False, indent=2)
    print(rendered)
    if args.json_out:
        args.json_out.parent.mkdir(parents=True, exist_ok=True)
        args.json_out.write_text(rendered + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
