from __future__ import annotations

import base64
from io import BytesIO
from typing import Any
from uuid import uuid4

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from fastapi.testclient import TestClient

from backend.chunk_cleaning import DocumentChunker
from backend import app as app_module


def _pdf_bytes(pages: list[list[str]]) -> bytes:
    writer = PdfWriter()
    for lines in pages:
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        font_ref = writer._add_object(font)
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})
        })
        commands = ["BT", "/F1 12 Tf", "72 720 Td"]
        for index, line in enumerate(lines):
            if index:
                commands.append("0 -24 Td")
            line = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            commands.append(f"({line}) Tj")
        commands.append("ET")
        stream = DecodedStreamObject()
        stream.set_data("\n".join(commands).encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def _cjk_pdf_bytes(pages: list[list[str]]) -> bytes:
    import fitz

    source = fitz.open()
    font_path = r"C:\Windows\Fonts\simhei.ttf"
    for lines in pages:
        page = source.new_page(width=612, height=792)
        for index, line in enumerate(lines):
            page.insert_text((72, 72 + index * 24), line, fontfile=font_path, fontname="cjk", fontsize=12)
    return source.tobytes()


def _rotated_pdf_bytes(rotation: int, pages: list[list[str]] | None = None) -> bytes:
    source = PdfReader(BytesIO(_pdf_bytes(pages or [["Rotated evidence"]])))
    writer = PdfWriter()
    page = source.pages[0]
    page.rotate(rotation)
    writer.add_page(page)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def _docx_bytes(token: str) -> bytes:
    document = Document()
    document.add_heading(f"Heading {token}", level=1)
    document.add_paragraph(f"中文说明 {token} with searchable details.")
    document.add_paragraph(f"List item {token}", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = f"Table header {token}"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = f"Table row {token}"
    table.cell(1, 1).text = "42"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _docx_with_footer_table() -> bytes:
    document = Document()
    footer = document.sections[0].footer
    footer.paragraphs[0].text = "Footer prose"
    table = footer.add_table(rows=2, cols=2, width=3000000)
    table.cell(0, 0).text = "Footer key"
    table.cell(0, 1).text = "Footer value"
    table.cell(1, 0).text = "Footer row"
    table.cell(1, 1).text = "7"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _docx_with_embedded_image() -> bytes:
    document = Document()
    document.add_heading("Image appendix", level=1)
    document.add_picture(BytesIO(_png_bytes()), width=None)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _xlsx_bytes(token: str) -> bytes:
    workbook = Workbook()
    first = workbook.active
    first.title = "Visible"
    first.append([f"Name {token}", "Amount"])
    first.append([f"Alpha {token}", 2])
    first.append([f"Beta {token}", 3])
    first["B4"] = "=SUM(B2:B3)"
    first.merge_cells("A5:B5")
    first["A5"] = f"Merged note {token}"
    second = workbook.create_sheet("Another")
    second.append([f"Second sheet {token}", "Value"])
    second.append([f"Second row {token}", 9])
    hidden = workbook.create_sheet("Hidden")
    hidden.append([f"Hidden sheet {token}", "Value"])
    hidden.append([f"Hidden row {token}", 7])
    hidden.sheet_state = "hidden"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _png_bytes() -> bytes:
    output = BytesIO()
    Image.new("RGB", (32, 20), color=(245, 245, 245)).save(output, format="PNG")
    return output.getvalue()


def _chunk_dicts(parsed: Any) -> list[dict[str, Any]]:
    assert hasattr(parsed, "chunks")
    assert hasattr(parsed, "stats")
    assert hasattr(parsed, "warnings")
    assert callable(parsed.to_dict)
    serialized = parsed.to_dict()
    assert isinstance(serialized, dict)
    assert {"chunks", "stats", "warnings"} <= serialized.keys()
    assert isinstance(serialized["chunks"], list)
    required = {
        "chunk_id", "document_id", "modality", "content_type", "text",
        "title_path", "page_no", "sheet_name", "row_start", "row_end",
        "source_locator", "parent_chunk_id", "quality", "warnings", "metadata",
    }
    chunks = []
    for chunk in parsed.chunks:
        assert callable(chunk.to_dict)
        record = chunk.to_dict()
        assert isinstance(record, dict)
        assert required <= record.keys()
        assert record["chunk_id"] and record["document_id"]
        assert record["modality"] and record["content_type"]
        assert isinstance(record["text"], str)
        assert isinstance(record["metadata"], dict)
        chunks.append(record)
    assert len(serialized["chunks"]) == len(parsed.chunks)
    return chunks


def _joined_text(chunks: list[dict[str, Any]]) -> str:
    return "\n".join(chunk["text"] for chunk in chunks)


def _rejected_or_warned(call: Any) -> None:
    try:
        parsed = call()
    except Exception:
        return
    assert not _chunk_dicts(parsed)
    assert parsed.warnings


def test_clean_text_preserves_unicode_and_is_idempotent() -> None:
    chunker = DocumentChunker(max_chars=1800, overlap_chars=120)
    source = "  中文标题\t\t和 日本語\r\n  café  e\u0301  😀  "
    cleaned = chunker.clean_text(source)
    assert "中文标题" in cleaned
    assert "日本語" in cleaned
    assert "café" in cleaned
    assert "😀" in cleaned
    assert chunker.clean_text(cleaned) == cleaned


def test_clean_text_handles_empty_and_whitespace_only_input() -> None:
    chunker = DocumentChunker(max_chars=1800, overlap_chars=120)
    assert chunker.clean_text("") == ""
    assert chunker.clean_text(" \t\r\n\u3000 ").strip() == ""


def test_pdf_keeps_page_locations_and_filters_repeated_page_boilerplate() -> None:
    token = uuid4().hex[:10]
    header, footer = f"Generated Header {token}", f"Generated Footer {token}"
    pdf = _pdf_bytes([
        [header, f"First page body {token}", footer],
        [header, f"Second page body {token}", footer],
    ])
    document_id = f"pdf-{token}"
    chunks = _chunk_dicts(DocumentChunker().parse_pdf(pdf, document_id=document_id))
    assert len(chunks) >= 2
    text = _joined_text(chunks)
    assert f"First page body {token}" in text
    assert f"Second page body {token}" in text
    assert header not in text and footer not in text
    assert {chunk["page_no"] for chunk in chunks} >= {1, 2}
    assert all(chunk["source_locator"] for chunk in chunks)
    assert all(chunk["document_id"] == document_id for chunk in chunks)


def test_pdf_cjk_headings_and_aligned_table_are_structured_with_page_context() -> None:
    pdf = _cjk_pdf_bytes([
        ["内部控制规范", "第一章 总则", "适用范围", "本制度适用于集团各部门"],
        ["内部控制规范", "|指标|阈值|", "|风险等级|3|", "第二章 管理", "异常事项应及时上报"],
    ])
    chunks = _chunk_dicts(DocumentChunker().parse_pdf(pdf, document_id="cjk-pdf"))
    assert any("第一章 总则" in chunk["text"] for chunk in chunks)
    assert any("第一章 总则" in chunk["title_path"] for chunk in chunks if chunk["page_no"] == 1)
    assert any(chunk["content_type"] == "row" and chunk["page_no"] == 2 for chunk in chunks)
    assert any("第二章 管理" in chunk["title_path"] for chunk in chunks if "异常事项" in chunk["text"])
    assert all(chunk["source_locator"] for chunk in chunks)


def test_pdf_cross_page_table_reuses_header_without_dropping_first_data_row() -> None:
    pdf = _cjk_pdf_bytes([
        ["|指标|阈值|"],
        ["|风险等级|3|"],
    ])
    chunks = _chunk_dicts(DocumentChunker().parse_pdf(pdf, document_id="cross-page-table"))
    rows = [chunk for chunk in chunks if chunk["content_type"] == "row"]
    assert len(rows) == 1
    assert rows[0]["page_no"] == 2
    assert "风险等级" in rows[0]["text"]


def test_pdf_adjacent_independent_tables_do_not_reuse_previous_header() -> None:
    pdf = _cjk_pdf_bytes([
        ["|指标|阈值|", "|风险等级|3|"],
        ["|部门|负责人|", "|研发|张三|"],
    ])
    chunks = _chunk_dicts(DocumentChunker().parse_pdf(pdf, document_id="independent-pdf-tables"))
    rows = [chunk for chunk in chunks if chunk["content_type"] == "row"]
    second_page_rows = [chunk for chunk in rows if chunk["page_no"] == 2]
    assert second_page_rows
    assert any("部门" in chunk["text"] and "负责人" in chunk["text"] for chunk in second_page_rows)
    assert all("指标" not in chunk["text"] for chunk in second_page_rows)


def test_pdf_chunk_ids_are_stable_for_same_document_and_content() -> None:
    token = uuid4().hex[:10]
    pdf = _pdf_bytes([[f"Stable PDF content {token}"]])
    chunker = DocumentChunker()
    first = _chunk_dicts(chunker.parse_pdf(pdf, document_id=f"stable-{token}"))
    second = _chunk_dicts(chunker.parse_pdf(pdf, document_id=f"stable-{token}"))
    ids = [chunk["chunk_id"] for chunk in first]
    assert ids == [chunk["chunk_id"] for chunk in second]
    assert len(set(ids)) == len(ids)


@pytest.mark.parametrize("rotation", [0, 90, 180, 270])
def test_pdf_rotation_coordinate_chain_is_complete_for_text_pages(rotation: int) -> None:
    parsed = DocumentChunker().parse_pdf(
        _rotated_pdf_bytes(rotation),
        document_id=f"rotated-{rotation}",
    )
    chunks = _chunk_dicts(parsed)
    text_chunk = next(chunk for chunk in chunks if "Rotated evidence" in chunk["text"])
    evidence = text_chunk["metadata"]["coordinate_evidence"]
    geometry = evidence["page_geometry"]
    assert geometry["rotation_degrees"] == rotation
    assert geometry["mapping_status"] == "complete"
    assert geometry["mappings"]["pdf_user_to_fitz_unrotated"]
    assert geometry["mappings"]["fitz_unrotated_to_display"]
    assert set(geometry["coordinate_systems"]) == {
        "pdf_user_space",
        "fitz_unrotated_text_space",
        "fitz_display_space",
        "render_pixel_space",
    }
    expected_size = [792.0, 612.0] if rotation in {90, 270} else [612.0, 792.0]
    assert geometry["display_size_pt"] == expected_size
    assert evidence["source_bbox_pdf_user_pt"]
    assert evidence["source_bbox_fitz_display_pt"]
    assert evidence["bbox_status"] == "exact_block"
    if rotation == 90:
        source_bbox = evidence["source_bbox_fitz_unrotated_pt"]
        display_bbox = evidence["source_bbox_fitz_display_pt"]
        assert display_bbox == pytest.approx(
            [792.0 - source_bbox[3], source_bbox[0], 792.0 - source_bbox[1], source_bbox[2]],
            abs=0.01,
        )
    assert parsed.stats["rotated_pages"] == ([1] if rotation else [])


def test_rotated_pdf_ocr_evidence_records_render_mapping() -> None:
    class RotatedPageOCR:
        def run(self, image_bytes: bytes, language: str = "eng") -> dict[str, Any]:
            assert image_bytes.startswith(b"\x89PNG")
            return {
                "status": "ok",
                "text": "旋转扫描页证据",
                "confidence": 0.91,
                "executor": "fixture",
                "attempts": [{"status": "accepted"}],
                "selected_transforms": ["exif_transpose"],
            }

    parsed = DocumentChunker().parse_pdf(
        _rotated_pdf_bytes(90, pages=[[]]),
        document_id="rotated-ocr",
        ocr_pipeline=RotatedPageOCR(),
        language="chi_sim+eng",
    )
    ocr_chunk = next(chunk for chunk in _chunk_dicts(parsed) if "旋转扫描页证据" in chunk["text"])
    evidence = ocr_chunk["metadata"]["coordinate_evidence"]
    geometry = evidence["page_geometry"]
    render = geometry["render"]
    assert geometry["rotation_degrees"] == 90
    assert render["rotation_applied"] is True
    assert render["width_px"] == 1584
    assert render["height_px"] == 1224
    assert evidence["render_input_sha256"]
    assert evidence["render_size_px"] == [1584, 1224]
    assert evidence["render_scale"] == 2.0
    assert evidence["bbox_status"] == "page_level"
    assert evidence["source_bbox_render_px"] == [0.0, 0.0, 1584.0, 1224.0]


def test_pdf_scanned_page_uses_ocr_pipeline_and_keeps_attempt_evidence() -> None:
    class PageOCR:
        def __init__(self) -> None:
            self.calls: list[str] = []

        def run(self, image_bytes: bytes, language: str = "eng") -> dict[str, Any]:
            assert image_bytes.startswith(b"\x89PNG")
            self.calls.append(language)
            return {
                "status": "ok",
                "text": "Recognized scanned page text with enough content.",
                "confidence": 0.91,
                "attempts": [{"attempt": 1, "status": "accepted"}],
                "selected_transforms": ["grayscale_normalize"],
            }

    ocr = PageOCR()
    parsed = DocumentChunker().parse_pdf(
        _pdf_bytes([[]]), document_id="scanned-page", ocr_pipeline=ocr, language="chi_sim+eng"
    )
    chunks = _chunk_dicts(parsed)
    assert ocr.calls == ["chi_sim+eng"]
    assert parsed.stats["ocr_succeeded_pages"] == [1]
    assert "Recognized scanned page text" in _joined_text(chunks)
    assert chunks[0]["quality"] == 0.91
    assert chunks[0]["metadata"]["ocr_attempts"][0]["status"] == "accepted"


def test_pdf_scanned_table_rows_keep_ocr_evidence() -> None:
    class TableOCR:
        def run(self, image_bytes: bytes, language: str = "eng") -> dict[str, Any]:
            return {
                "status": "ok",
                "text": "|指标|阈值|\n|风险等级|3|",
                "confidence": 0.88,
                "executor": "fixture",
                "attempts": [{"status": "accepted"}],
                "selected_transforms": ["grayscale_normalize"],
            }

    parsed = DocumentChunker().parse_pdf(
        _pdf_bytes([[]]),
        document_id="scanned-table-evidence",
        ocr_pipeline=TableOCR(),
        language="chi_sim+eng",
    )
    rows = [chunk.to_dict() for chunk in parsed.chunks if chunk.content_type == "row"]
    assert rows
    evidence = rows[0]["metadata"]["ocr_evidence"]
    assert evidence["executor"] == "fixture"
    assert evidence["language"] == "chi_sim+eng"
    assert evidence["confidence"] == 0.88
    assert evidence["attempts"][0]["status"] == "accepted"
    assert evidence["bbox_status"] == "unavailable"


def test_image_content_hash_changes_stable_chunk_id() -> None:
    first = _png_bytes()
    output = BytesIO()
    Image.new("RGB", (32, 20), color=(230, 230, 230)).save(output, format="PNG")
    second = output.getvalue()
    first_chunk = _chunk_dicts(DocumentChunker().parse_image(first, document_id="same-image"))[0]
    second_chunk = _chunk_dicts(DocumentChunker().parse_image(second, document_id="same-image"))[0]
    assert first_chunk["metadata"]["sha256"] != second_chunk["metadata"]["sha256"]
    assert first_chunk["chunk_id"] != second_chunk["chunk_id"]


def test_docx_preserves_heading_list_table_and_source_traceability() -> None:
    token = uuid4().hex[:10]
    document_id = f"docx-{token}"
    parsed = DocumentChunker().parse_docx(_docx_bytes(token), document_id=document_id)
    chunks = _chunk_dicts(parsed)
    text = _joined_text(chunks)
    assert chunks
    assert f"Heading {token}" in text
    assert f"中文说明 {token}" in text
    assert f"List item {token}" in text
    assert f"Table header {token}" in text and f"Table row {token}" in text
    assert all(chunk["source_locator"] for chunk in chunks)
    assert all(chunk["document_id"] == document_id for chunk in chunks)
    assert any(chunk["title_path"] for chunk in chunks)


def test_docx_header_footer_table_is_not_dropped() -> None:
    chunks = _chunk_dicts(DocumentChunker().parse_docx(_docx_with_footer_table(), document_id="footer-table"))
    text = _joined_text(chunks)
    assert "Footer key" in text and "Footer row" in text
    assert any(chunk["metadata"].get("header_footer") == "footer" for chunk in chunks)


def test_docx_embedded_image_can_produce_auditable_ocr_child_chunk() -> None:
    class EmbeddedImageOCR:
        def __init__(self) -> None:
            self.languages: list[str] = []

        def run(self, image_bytes: bytes, language: str = "eng") -> dict[str, Any]:
            assert image_bytes.startswith(b"\x89PNG")
            self.languages.append(language)
            return {
                "status": "ok",
                "text": "印章页审批编号 A-17",
                "confidence": 0.93,
                "attempts": [{"status": "accepted"}],
                "selected_transforms": ["grayscale_normalize"],
                "executor": "fixture",
            }

    ocr = EmbeddedImageOCR()
    parsed = DocumentChunker().parse_docx(
        _docx_with_embedded_image(),
        document_id="docx-embedded-image",
        ocr_pipeline=ocr,
        language="chi_sim+eng",
    )
    chunks = _chunk_dicts(parsed)
    image_chunks = [chunk for chunk in chunks if chunk["content_type"] == "image"]
    ocr_chunks = [chunk for chunk in chunks if chunk["content_type"] == "ocr_region"]
    assert ocr.languages == ["chi_sim+eng"]
    assert len(image_chunks) == 1
    assert image_chunks[0]["metadata"]["sha256"]
    assert image_chunks[0]["metadata"]["ocr_status"] == "ok"
    assert len(ocr_chunks) == 1
    assert "审批编号 A-17" in ocr_chunks[0]["text"]
    assert ocr_chunks[0]["parent_chunk_id"] == image_chunks[0]["chunk_id"]
    assert ocr_chunks[0]["metadata"]["ocr_executor"] == "fixture"


def test_xlsx_preserves_sheets_formulas_hidden_data_and_merged_cells() -> None:
    token = uuid4().hex[:10]
    document_id = f"xlsx-{token}"
    parsed = DocumentChunker().parse_xlsx(_xlsx_bytes(token), document_id=document_id)
    chunks = _chunk_dicts(parsed)
    text = _joined_text(chunks)
    assert chunks
    for expected in (
        f"Alpha {token}", f"Second sheet {token}", f"Second row {token}",
        f"Merged note {token}",
    ):
        assert expected in text
    assert f"Hidden sheet {token}" not in text
    assert f"Hidden row {token}" not in text
    assert any("hidden_sheet_skipped:Hidden" in warning for warning in parsed.warnings)
    assert any(
        "=SUM(B2:B3)" in chunk["text"]
        or "=SUM(B2:B3)" in str(chunk["metadata"])
        for chunk in chunks
    )
    assert {chunk["sheet_name"] for chunk in chunks} >= {"Visible", "Another"}
    assert "Hidden" not in {chunk["sheet_name"] for chunk in chunks}
    assert all(chunk["source_locator"] for chunk in chunks)
    assert all(chunk["document_id"] == document_id for chunk in chunks)
    assert any(chunk["row_start"] is not None and chunk["row_end"] is not None for chunk in chunks)


class _LowConfidenceOCR:
    def __init__(self, confidence: float) -> None:
        self.confidence = confidence
        self.languages: list[str] = []

    def run(self, image_bytes: bytes, language: str = "eng") -> dict[str, Any]:
        assert image_bytes
        self.languages.append(language)
        return {
            "status": "ok",
            "text": "OCR extracted text",
            "confidence": self.confidence,
        }


def test_image_records_low_ocr_confidence_and_requested_language() -> None:
    token = uuid4().hex[:10]
    ocr = _LowConfidenceOCR(confidence=0.2)
    parsed = DocumentChunker().parse_image(
        _png_bytes(), document_id=f"image-{token}", ocr_pipeline=ocr, language="jpn"
    )
    chunks = _chunk_dicts(parsed)
    assert ocr.languages == ["jpn"]
    assert chunks
    assert "OCR extracted text" in _joined_text(chunks)
    assert all(chunk["source_locator"] for chunk in chunks)
    assert all(chunk["document_id"] == f"image-{token}" for chunk in chunks)
    diagnostics = str(parsed.warnings) + str([chunk["metadata"] for chunk in chunks])
    assert "0.2" in diagnostics or "confidence" in diagnostics.lower()


def test_image_ocr_failure_is_reported_without_losing_parse_result() -> None:
    class FailingOCR:
        def run(self, image_bytes: bytes, language: str = "eng") -> Any:
            raise RuntimeError("generated OCR failure")

    parsed = DocumentChunker().parse_image(
        _png_bytes(), document_id="image-ocr-failure", ocr_pipeline=FailingOCR()
    )
    _chunk_dicts(parsed)
    diagnostics = str(parsed.warnings) + str(parsed.stats)
    assert "OCR" in diagnostics.upper()
    assert "generated OCR failure" in diagnostics or parsed.warnings


def test_parse_methods_return_structured_empty_results_for_empty_documents() -> None:
    token = uuid4().hex[:10]
    empty_pdf = _pdf_bytes([])
    empty_docx = BytesIO()
    Document().save(empty_docx)
    empty_xlsx = BytesIO()
    Workbook().save(empty_xlsx)
    results = [
        DocumentChunker().parse_pdf(empty_pdf, document_id=f"empty-pdf-{token}"),
        DocumentChunker().parse_docx(empty_docx.getvalue(), document_id=f"empty-docx-{token}"),
        DocumentChunker().parse_xlsx(empty_xlsx.getvalue(), document_id=f"empty-xlsx-{token}"),
    ]
    for parsed in results:
        assert not _chunk_dicts(parsed)


@pytest.mark.parametrize(
    ("method_name", "payload"),
    [("parse_pdf", b"not a PDF"), ("parse_docx", b"not a DOCX"), ("parse_xlsx", b"not an XLSX")],
)
def test_malformed_office_documents_are_rejected_or_return_a_warning(
    method_name: str, payload: bytes
) -> None:
    method = getattr(DocumentChunker(), method_name)
    _rejected_or_warned(lambda: method(payload, document_id="invalid-input"))


def test_malformed_image_is_rejected_or_return_a_warning() -> None:
    _rejected_or_warned(
        lambda: DocumentChunker().parse_image(
            b"not an image",
            document_id="invalid-image",
            ocr_pipeline=_LowConfidenceOCR(0.9),
        )
    )


def test_large_paragraph_splits_with_bounded_overlap_without_truncation() -> None:
    chunker = DocumentChunker(max_chars=160, overlap_chars=20)
    source = ("多模态检索证据需要可追溯；enterprise provenance matters. " * 30).strip()
    chunks = chunker._make_chunks(
        "large-paragraph", "docx", "paragraph", source,
        locator="body:paragraph:1", overlap=True,
    )
    assert len(chunks) > 1
    assert all(len(chunk.text) <= 160 for chunk in chunks)
    for phrase in ("多模态检索证据", "enterprise provenance matters"):
        assert all(phrase not in chunk.text for chunk in chunks) is False
    assert all(chunk.source_locator for chunk in chunks)


def test_xlsx_does_not_drop_first_row_when_header_is_ambiguous() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append([101, 202])
    worksheet.append([303, 404])
    output = BytesIO()
    workbook.save(output)
    chunks = _chunk_dicts(DocumentChunker().parse_xlsx(output.getvalue(), document_id="numeric-rows"))
    text = _joined_text(chunks)
    assert "101" in text and "202" in text
    assert "303" in text and "404" in text


def test_xlsx_repeated_headers_and_numeric_rows_are_conservative() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append([101, 202])
    worksheet.append([303, 404])
    worksheet.append([101, 202])
    output = BytesIO()
    workbook.save(output)
    parsed = DocumentChunker().parse_xlsx(output.getvalue(), document_id="ambiguous-header")
    text = _joined_text(_chunk_dicts(parsed))
    assert "101" in text and "202" in text and "303" in text and "404" in text
    assert "header_ambiguous" in " ".join(parsed.warnings)


def test_xlsx_skips_hidden_rows_and_columns_and_marks_formula_cache_unverified() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sensitive"
    worksheet.append(["Name", "Amount", "Internal note"])
    worksheet.append(["Visible row", 3, "visible detail"])
    worksheet.append(["Hidden row", 4, "secret row"])
    worksheet["B4"] = "=SUM(B2:B3)"
    worksheet.row_dimensions[3].hidden = True
    worksheet.column_dimensions["C"].hidden = True
    output = BytesIO()
    workbook.save(output)

    parsed = DocumentChunker().parse_xlsx(output.getvalue(), document_id="xlsx-visibility")
    chunks = _chunk_dicts(parsed)
    text = _joined_text(chunks)
    assert "Visible row" in text
    assert "Hidden row" not in text
    assert "secret row" not in text
    formula_chunks = [chunk for chunk in chunks if "SUM(B2:B3)" in str(chunk["metadata"])]
    assert formula_chunks
    formula_values = formula_chunks[0]["metadata"]["values"]
    assert any(value.get("cache_status") == "missing" for value in formula_values if isinstance(value, dict))
    assert any("hidden_rows_skipped:Sensitive" in warning for warning in parsed.warnings)
    assert any("hidden_columns_skipped:Sensitive" in warning for warning in parsed.warnings)


def test_chunk_preview_endpoint_returns_provenance() -> None:
    token = uuid4().hex[:10]
    response = TestClient(app_module.app).post(
        "/api/v1/documents/chunks-preview",
        json={
            "document_id": f"api-{token}",
            "modality": "pdf",
            "file_base64": base64.b64encode(_pdf_bytes([[f"API source {token}"]])).decode("ascii"),
        },
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["document_id"] == f"api-{token}"
    assert payload["chunks"]
    assert payload["chunks"][0]["source_locator"].startswith("page:1:")
    assert f"API source {token}" in _joined_text(payload["chunks"])
