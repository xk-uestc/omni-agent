"""Independent source/typed-cell regressions; these are not model scores."""
from copy import deepcopy
from io import BytesIO
import json
import sqlite3

import pytest
from openpyxl import Workbook

from backend.dependency_agent import DependencyAgent
from backend.fusion_constraints import (
    SourceConstraintError, VerifiedFormulaTarget, bind_source_constraints,
    verify_required_intent,
)
from backend.knowledge_store import KnowledgeStore
from backend.nl2sql.engine import Nl2SqlEngine
from backend.omni_agent import OmniAgent, _check_static_document_cells
from backend.responses_client import GenerationError
from backend.session import ConversationStore


@pytest.fixture
def environment(tmp_path):
    database = tmp_path / 'parcels.sqlite'
    with sqlite3.connect(database) as connection:
        connection.execute('CREATE TABLE 包裹台账(包裹ID INTEGER PRIMARY KEY, 地区 TEXT, 出库日期 TEXT, 费用金额 REAL)')
        connection.executemany('INSERT INTO 包裹台账 VALUES(?,?,?,?)', [
            (1, '南湾', '2025-03-01', 40), (2, '南湾', '2025-05-01', 20),
            (3, '南湾', '2024-03-01', 200), (4, '北岭', '2025-04-01', 100),
        ])
    engine = Nl2SqlEngine(database, metric_catalog_path=tmp_path / 'absent.json')
    knowledge = KnowledgeStore(tmp_path / 'knowledge')
    ingest_table(knowledge, [('南湾', 2026, .15), ('北岭', 2026, .10)])
    return engine, knowledge


def ingest_table(knowledge, rows):
    book = Workbook()
    book.active.append(['地区', '年份', '调增率'])
    for values in rows:
        book.active.append(values)
        book.active.cell(book.active.max_row, 3).number_format = '0.0%'
    stream = BytesIO()
    book.save(stream)
    knowledge.ingest(stream.getvalue(), document_id='freight-targets', title='运输目标',
                     modality='xlsx', filename='targets.xlsx')


SOURCE_QUESTION = '根据计算文档净核算额公式，以及数据库2025年南湾包裹台账费用金额合计，算2026年净核算额。'
CELL_QUESTION = '从运输目标Excel中定位2026年调增率为15%的地区，作为过滤条件查询数据库中该地区2025年的费用金额合计。'


def sql_tasks(question='2025年南湾包裹台账费用金额合计'):
    return [{'id': 'amount', 'tool': 'sql', 'args': {'question': question}}]


def cell_tasks(value=.15, *, extra_where=None):
    return [
        {'id': 'region', 'tool': 'document_cell', 'args': {
            'document_id': 'freight-targets', 'where': {'年份': 2026, '调增率': value, **(extra_where or {})},
            'column': '地区'}},
        {'id': 'amount', 'tool': 'sql', 'args': {'question': [
            '2025年包裹台账费用金额合计，地区为', {'ref': 'region', 'path': ['value']}]}},
    ]


def test_source_connector_keeps_baseline_and_separate_formula_target(environment):
    engine, _ = environment
    proof = VerifiedFormulaTarget('净核算额', 'local-formula', 'a' * 64, 'formula')
    bound, audit = bind_source_constraints(SOURCE_QUESTION, sql_tasks(), engine,
                                           verified_formula_targets=(proof,))
    required = bound['amount']
    assert verify_required_intent(engine.extract_required_intent(sql_tasks()[0]['args']['question']), required) == []
    assert audit[0]['text'] == '2025年南湾包裹台账费用金额合计'
    assert audit[0]['target_binding']['target_year'] == 2026
    assert audit[0]['target_binding']['label'] == '净核算额'
    assert '2026年' not in audit[0]['text']
    assert 'source_filter_mismatch' in verify_required_intent(
        engine.extract_required_intent('2026年南湾包裹台账费用金额合计'), required)
    assert 'source_metric_mismatch' in verify_required_intent(
        engine.extract_required_intent('2025年南湾包裹台账费用金额平均'), required)


@pytest.mark.parametrize('prefix', ['已批准', '南湾', '不含南湾', '核验后记录才参与'])
def test_connector_never_consumes_unknown_or_unbound_business_scope(environment, prefix):
    engine, _ = environment
    proof = VerifiedFormulaTarget('净核算额', 'local-formula', 'a' * 64, 'formula')
    question = SOURCE_QUESTION.replace('以及数据库', f'以及{prefix}数据库')
    with pytest.raises(SourceConstraintError, match='来源范围'):
        bind_source_constraints(question, sql_tasks(), engine, verified_formula_targets=(proof,))


def test_cell_preview_preserves_json_value_and_execution_unit(environment):
    engine, knowledge = environment
    catalogue = OmniAgent(engine, knowledge, ConversationStore()).catalogue(CELL_QUESTION)
    document = next(item for item in catalogue if item['id'] == 'freight-targets')
    assert '15%' in document['text']
    row = next(item for item in document['table_rows'] if item['values'][0] == '南湾')
    index = row['headers'].index('调增率')
    assert type(row['values'][index]) is float and row['values'][index] == .15
    assert type(row['values'][row['headers'].index('年份')]) is int
    evidence = DependencyAgent(engine, knowledge).execute('document_cell', {
        'document_id': 'freight-targets', 'where': {'地区': '南湾', '年份': 2026}, 'column': '调增率'}, {}, {})
    assert evidence['value'] == .15 and evidence['unit'] == 'ratio'


@pytest.mark.parametrize('value', ['15%', '0.15', 15, .16])
def test_literal_cell_preflight_rejects_wrong_representation_without_conversion(environment, value):
    _, knowledge = environment
    tasks = cell_tasks(value)
    original = deepcopy(tasks)
    with pytest.raises(GenerationError) as error:
        _check_static_document_cells(tasks, knowledge)
    assert error.value.plan_rejection_code == 'document_cell_static_selection_unverified'
    assert tasks == original


def test_literal_cell_preflight_reads_full_source_even_when_preview_omits_match(environment, monkeypatch):
    engine, knowledge = environment
    ingest_table(knowledge, [(f'其他{i}', 2026, .10) for i in range(35)] + [('南湾', 2026, .15)])
    monkeypatch.setattr(knowledge, 'search', lambda *args, **kwargs: [])
    agent = OmniAgent(engine, knowledge, ConversationStore())
    preview = agent.catalogue(CELL_QUESTION)
    document = next(item for item in preview if item['id'] == 'freight-targets')
    assert len(document['table_rows']) == 30 and document['omitted_table_rows'] == 6
    assert all(row['values'][0] != '南湾' for row in document['table_rows'])
    _check_static_document_cells(cell_tasks(), knowledge)
    agent.client = ReplayPlanner([cell_tasks()])
    result = agent.query(CELL_QUESTION)
    assert result['status'] == 'ok' and len(agent.client.calls) == 1
    assert result['result']['results']['amount']['rows'][0]['总费用金额'] == 60


def test_literal_cell_preflight_rejects_duplicate_rows_without_deduplication(environment):
    _, knowledge = environment
    ingest_table(knowledge, [('南湾', 2026, .15), ('南湾', 2026, .15)])
    with pytest.raises(GenerationError) as error:
        _check_static_document_cells(cell_tasks(), knowledge)
    assert error.value.plan_rejection_code == 'document_cell_static_selection_unverified'


@pytest.mark.parametrize('condition', [{'年份': '2026'}, {'未给定业务条件': '已批准'}])
def test_literal_cell_preflight_keeps_unknown_conditions_and_scalar_types(environment, condition):
    _, knowledge = environment
    tasks = cell_tasks(extra_where=condition)
    original = deepcopy(tasks)
    with pytest.raises(GenerationError) as error:
        _check_static_document_cells(tasks, knowledge)
    assert error.value.plan_rejection_code == 'document_cell_static_selection_unverified'
    assert tasks == original


def test_dynamic_cell_arguments_remain_execution_time_obligations(environment, monkeypatch):
    _, knowledge = environment
    tasks = cell_tasks()
    tasks[0]['args']['where']['年份'] = {'ref': 'upstream', 'path': ['value']}
    monkeypatch.setattr(knowledge, 'document', lambda identifier: pytest.fail('Dynamic cell was pre-executed'))
    _check_static_document_cells(tasks, knowledge)


class ReplayPlanner:
    audit = {'status': 'completed', 'http_status': 200}

    def __init__(self, responses):
        self.responses, self.calls = responses, []

    def generate(self, instructions, context, schema, **kwargs):
        self.calls.append({'context': deepcopy(context), 'operation': kwargs['name'], 'instructions': instructions})
        tasks = self.responses[min(len(self.calls) - 1, len(self.responses) - 1)]
        return {'route': 'fusion', 'effective_question': context['question'], 'clarification': '',
                'tasks_json': json.dumps(tasks, ensure_ascii=False)}


def test_invalid_literal_gets_one_replan_before_tools_and_keeps_corrected_plan(environment, monkeypatch):
    engine, knowledge = environment
    wrong, corrected = cell_tasks('15%'), cell_tasks()
    planner = ReplayPlanner([wrong, corrected])
    executed = []
    original = DependencyAgent.execute

    def observe(self, tool, args, *rest, **kwargs):
        assert len(planner.calls) == 2
        executed.append(tool)
        return original(self, tool, args, *rest, **kwargs)

    monkeypatch.setattr(DependencyAgent, 'execute', observe)
    result = OmniAgent(engine, knowledge, ConversationStore(), planner).query(CELL_QUESTION)
    assert result['status'] == 'ok' and result['planner_source'] == 'model_validated'
    assert result['result']['results']['amount']['rows'][0]['总费用金额'] == 60
    assert len(planner.calls) == 2 and planner.calls[1]['operation'] == 'omni_plan_completion_repair'
    assert planner.calls[1]['context']['plan_completion_feedback']['errors'] == ['document_cell_static_selection_unverified']
    assert result['result']['execution_plan'] == corrected
    assert wrong[0]['args']['where']['调增率'] == '15%'
    assert [attempt['validation'] for attempt in result['trace'][0]['attempts']] == ['plan_protocol_rejected', 'complete']
    assert executed and set(executed) == {'document_cell', 'sql'}


def test_repeated_invalid_literal_stops_after_one_repair_without_tools(environment, monkeypatch):
    engine, knowledge = environment
    planner = ReplayPlanner([cell_tasks('15%')])
    monkeypatch.setattr(DependencyAgent, 'execute', lambda *args, **kwargs: pytest.fail('Rejected literal executed'))
    result = OmniAgent(engine, knowledge, ConversationStore(), planner).query(CELL_QUESTION)
    assert len(planner.calls) == 2 and result['status'] == 'clarification'
    assert result['planner_source'] == 'rules_fallback'


def test_matching_extra_condition_still_cannot_authorize_user_filter(environment, monkeypatch):
    engine, knowledge = environment
    tasks = cell_tasks(extra_where={'地区': '南湾'})
    _check_static_document_cells(tasks, knowledge)  # Tool contract alone does not prove user intent.
    monkeypatch.setattr(engine, 'answer', lambda *args, **kwargs: pytest.fail('Added output filter reached SQL'))
    result = OmniAgent(engine, knowledge, ConversationStore(), ReplayPlanner([tasks])).query(CELL_QUESTION)
    assert result['status'] == 'incomplete'
    assert result['result']['error_code'] == 'source_constraint_mismatch'
    assert result['result']['trace'][-1]['source_constraint_code'] == 'source_dynamic_binding_unverified'


def test_static_cell_preflight_rejects_changed_or_unavailable_original(environment):
    _, knowledge = environment
    source, _ = knowledge.original('freight-targets')
    source.write_bytes(b'changed local fixture')
    with pytest.raises(GenerationError) as error:
        _check_static_document_cells(cell_tasks(), knowledge)
    assert error.value.plan_rejection_code == 'document_cell_static_source_unverified'
