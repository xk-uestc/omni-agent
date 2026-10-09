from io import BytesIO
import json

import pytest
from openpyxl import Workbook

from backend.knowledge_store import KnowledgeStore
from backend.nl2sql.engine import Nl2SqlEngine
from backend.nl2sql.seed import initialize_database
from backend.omni_agent import OmniAgent
from backend.session import ConversationStore
from backend.fusion_history import resolve_fusion_followup
from backend.fusion_constraints import SourceConstraintError
from backend.dependency_agent import DependencyAgent


FIRST = '按指标公式文档计算2025年华东地区客单价，销售额和订单数从数据库取。'
FORECAST = '根据经营预测PDF公式、区域目标Excel中华东2026增长率和数据库2025年华东销售额，算2026年目标销售额。'


def tasks_for(question):
    region = '华南' if '华南' in question else '华东'
    forecast = '目标销售额' in question
    tasks = [
        {'id': 'formula', 'tool': 'document_formula', 'args': {'document_id': 'policy', 'label': '目标销售额' if forecast else '客单价'}},
        {'id': 'sql', 'tool': 'sql', 'args': {'question': f'2025年{region}地区销售额' + ('' if forecast else '和订单数')}},
    ]
    parameters = {'销售额': {'ref': 'sql', 'path': ['rows', 0, '销售额']},
                  '订单数': {'ref': 'sql', 'path': ['rows', 0, '订单数']}}
    if forecast:
        tasks.append({'id': 'growth', 'tool': 'document_cell', 'args': {'document_id': 'targets', 'where': {'地区': region, '年份': 2026}, 'column': '目标增长率'}})
        parameters = {'基准销售额': {'ref': 'sql', 'path': ['rows', 0, '销售额']}, '目标增长率': {'ref': 'growth', 'path': []}}
    tasks.append({'id': 'calculate', 'tool': 'calculate', 'args': {'formula': {'ref': 'formula', 'path': []}, 'parameters': parameters}})
    return tasks


@pytest.fixture
def agent(tmp_path):
    knowledge = KnowledgeStore(tmp_path / 'knowledge')
    knowledge.ingest('客单价 = 销售额 / 订单数\n目标销售额 = 基准销售额 * (1 + 目标增长率)'.encode(), document_id='policy', title='指标与预测公式', modality='txt', filename='p.txt')
    workbook = Workbook()
    workbook.active.append(['地区', '年份', '目标增长率'])
    workbook.active.append(['华东', 2026, .12])
    workbook.active.append(['华南', 2026, .10])
    buffer = BytesIO()
    workbook.save(buffer)
    knowledge.ingest(buffer.getvalue(), document_id='targets', title='区域目标', modality='xlsx', filename='t.xlsx')
    class Planner:
        contexts = []
        def generate(self, instructions, context, *args, **kwargs):
            self.contexts.append(context)
            return {'route': 'fusion', 'effective_question': '2024年华北销售额', 'clarification': '',
                    'tasks_json': json.dumps(tasks_for(context['question']), ensure_ascii=False)}
    return OmniAgent(Nl2SqlEngine(initialize_database(tmp_path / 'sales.sqlite')), knowledge,
                     ConversationStore(storage_path=tmp_path / 'sessions.sqlite'), Planner())


def test_four_cross_source_turns_use_verified_server_scopes(agent):
    one = agent.query(FIRST, session_id='cross')
    assert one['status'] == 'ok', one
    assert one['effective_question'] == FIRST
    assert one['state']['fusion_context']['source_bindings']
    two = agent.query('那华南呢，仍按同一公式？', session_id='cross')
    assert two['status'] == 'ok', two
    assert two['result']['results']['calculate']['value'] == pytest.approx(22992 / 3)
    assert '2025年华南' in two['effective_question'] and '2024年' not in two['effective_question']
    assert two['context_resolution']['actual_question'] == two['question']
    three = agent.query(FORECAST, session_id='cross')
    assert three['status'] == 'ok', three
    assert three['context_resolution']['mode'] == 'independent'
    four = agent.query('那华南2026年的目标呢，基准还是2025年，增长率取该地区Excel？', session_id='cross')
    assert four['status'] == 'ok', four
    assert four['result']['results']['calculate']['value'] == pytest.approx(22992 * 1.10)
    assert '数据库2025年华南' in four['effective_question']
    assert '华南2026增长率' in four['effective_question']
    assert set(four['state']['fusion_context']['documents']) == {'policy', 'targets'}


def test_same_algorithm_reuses_sealed_tasks_with_fresh_region_reads(agent,monkeypatch):
    first=agent.query('先查2025年华东销售额作预测基准，再结合《指标与预测公式》的公式和《区域目标》中华东2026目标增长率，算出华东2026目标销售额，并分别标明PDF、Excel和数据库来源。',session_id='template')
    assert first['status']=='ok',first
    def forbidden(*args,**kwargs):pytest.fail('Verified region replacement must not call the model planner')
    monkeypatch.setattr(agent.client,'generate',forbidden)
    second=agent.query('同样算法改成华南，继续用2025销售额为基准，并使用区域目标对应增长率。',session_id='template')
    assert second['status']=='ok',second
    assert second['result']['results']['calculate']['value']==pytest.approx(22992*1.1)
    assert second['trace'][0]['source']=='server_verified_fusion_task_template'


def test_tampered_task_template_never_reaches_execution(agent,monkeypatch):
    assert agent.query(FORECAST,session_id='sealed')['status']=='ok'
    turn=agent.conversations.context('sealed')[-1]
    state=turn.state
    state['fusion_context']['task_template']['payload']['tasks'][2]['args']['where']['年份']=2025
    agent.conversations.remember('sealed',question=turn.question,effective_question=turn.effective_question,state=state)
    monkeypatch.setattr(DependencyAgent,'run',lambda *args,**kwargs:pytest.fail('Tampered task template reached execution'))
    result=agent.query('那华南2026年的目标呢，基准还是2025年，增长率取该地区Excel？',session_id='sealed')
    assert result['status']=='clarification'


@pytest.mark.parametrize('question', [
    '那华南2025年的目标呢，基准还是2026年，增长率取该地区Excel？',
    '那华南2027年的目标呢，基准还是2025年，增长率取该地区Excel？',
    '那2026年呢', '那华南呢，不含线上', '那华南和华北呢',
])
def test_ambiguous_time_or_unsupported_slot_clarifies_before_model(agent, question):
    assert agent.query(FORECAST, session_id='bad')['status'] == 'ok'
    calls = len(agent.client.contexts)
    result = agent.query(question, session_id='bad')
    assert result['status'] == 'clarification'
    assert len(agent.client.contexts) == calls
    assert not result['state'].get('fusion_context')


def test_failed_turn_cannot_inherit_earlier_success(agent):
    assert agent.query(FIRST, session_id='failed')['status'] == 'ok'
    assert agent.query('那华南呢，不含线上', session_id='failed')['status'] == 'clarification'
    result = agent.query('那华南呢，仍按同一公式？', session_id='failed')
    assert result['status'] != 'ok'
    assert result['context_resolution']['mode'] == 'independent'


def test_new_topic_and_reset_do_not_inherit_sources(agent):
    assert agent.query(FIRST, session_id='topic')['status'] == 'ok'
    result = agent.query('换个主题：2024年华北地区销售额', session_id='topic')
    assert result['context_resolution']['mode'] == 'independent'
    result = agent.query('那华南呢', session_id='topic', reset_context=True)
    assert result['context_turns'] == 0
    assert result['context_resolution']['mode'] == 'independent'


def test_model_cannot_switch_verified_document_in_followup(agent):
    assert agent.query(FIRST, session_id='source')['status'] == 'ok'
    # Legacy sessions predate sealed task templates; retain coverage of the
    # model replanning fallback as well as the new template path.
    turn=agent.conversations.context('source')[-1]
    state=dict(turn.state);state['fusion_context']=dict(state['fusion_context'])
    state['fusion_context'].pop('task_template',None)
    agent.conversations.remember('source',question=turn.question,effective_question=turn.effective_question,state=state)
    original = agent.client.generate
    def changed(*args, **kwargs):
        plan = original(*args, **kwargs)
        tasks = json.loads(plan['tasks_json'])
        tasks[0]['args']['document_id'] = 'targets'
        plan['tasks_json'] = json.dumps(tasks)
        return plan
    agent.client.generate = changed
    result = agent.query('那华南呢，仍按同一公式？', session_id='source')
    assert result['status'] == 'clarification'
    assert result['result']['clarification_code'] == 'fusion_followup_source_changed'
    assert result['result']['results'] == {}


@pytest.mark.parametrize('mutation', ['label', 'column', 'year'])
def test_same_documents_cannot_authorize_changed_parameter_contracts(agent, mutation, monkeypatch):
    assert agent.query(FORECAST, session_id='bindings')['status'] == 'ok'
    turn=agent.conversations.context('bindings')[-1]
    state=dict(turn.state);state['fusion_context']=dict(state['fusion_context'])
    state['fusion_context'].pop('task_template',None)
    agent.conversations.remember('bindings',question=turn.question,effective_question=turn.effective_question,state=state)
    monkeypatch.setattr(DependencyAgent, 'run',
                        lambda *args, **kwargs: pytest.fail('Changed parameter contract reached tool execution'))
    original = agent.client.generate
    def changed(*args, **kwargs):
        plan = original(*args, **kwargs)
        tasks = json.loads(plan['tasks_json'])
        if mutation == 'label':
            tasks[0]['args']['label'] = '客单价'
        elif mutation == 'column':
            tasks[2]['args']['column'] = '年份'
        else:
            tasks[2]['args']['where']['年份'] = 2025
        plan['tasks_json'] = json.dumps(tasks)
        return plan
    agent.client.generate = changed
    result = agent.query('那华南2026年的目标呢，基准还是2025年，增长率取该地区Excel？', session_id='bindings')
    assert result['status'] == 'clarification'
    assert not result['state'].get('fusion_context')
    if mutation == 'year':
        # No source row has that year: the literal tool contract now rejects
        # the plan before inherited binding validation or any DAG execution.
        assert result['trace'][0]['rejection_code'] == 'document_cell_static_selection_unverified'
        assert all(attempt['errors'] == ['document_cell_static_selection_unverified']
                   for attempt in result['trace'][0]['attempts'])
    else:
        assert result['result']['clarification_code'] == 'fusion_followup_document_binding_changed'
        assert result['result']['results'] == {}


def test_entity_substring_in_document_title_is_not_replaced(agent):
    assert agent.query(FORECAST, session_id='title')['status'] == 'ok'
    previous = agent.conversations.context('title')[-1]
    # A schema value embedded in a filename/title is not a scalar parameter.
    previous.state['fusion_context']['scope_question'] += '参考区域华东地区报告'
    with pytest.raises(SourceConstraintError, match='跨源查询') as error:
        resolve_fusion_followup('那华南2026年的目标呢，基准还是2025年，增长率取该地区Excel？',
                               [previous], agent.engine, agent.knowledge)
    assert error.value.code == 'fusion_followup_ambiguous'


def test_changed_document_revision_does_not_inherit(agent):
    assert agent.query(FORECAST, session_id='revision')['status'] == 'ok'
    agent.knowledge.ingest('目标销售额 = 基准销售额 * (1 + 目标增长率)\n变更版本'.encode(),
                           document_id='policy', title='变更', modality='txt', filename='new.txt')
    calls = len(agent.client.contexts)
    result = agent.query('那华南2026年的目标呢，基准还是2025年，增长率取该地区Excel？', session_id='revision')
    assert result['status'] == 'clarification'
    assert len(agent.client.contexts) == calls


def test_followup_cannot_drop_document_route(agent):
    assert agent.query(FORECAST, session_id='route')['status'] == 'ok'
    turn=agent.conversations.context('route')[-1]
    state=dict(turn.state);state['fusion_context']=dict(state['fusion_context'])
    state['fusion_context'].pop('task_template',None)
    agent.conversations.remember('route',question=turn.question,effective_question=turn.effective_question,state=state)
    def changed(*args, **kwargs):
        return {'route': 'sql', 'effective_question': '2025年华南销售额', 'tasks_json': '[]', 'clarification': ''}
    agent.client.generate = changed
    result = agent.query('那华南2026年的目标呢，基准还是2025年，增长率取该地区Excel？', session_id='route')
    assert result['status'] == 'clarification'
    assert result['route'] == 'clarify'


def test_verified_fusion_context_survives_session_store_restart(agent):
    assert agent.query(FORECAST, session_id='persist')['status'] == 'ok'
    reopened = OmniAgent(agent.engine, agent.knowledge,
                         ConversationStore(storage_path=agent.conversations.storage_path), agent.client)
    result = reopened.query('那华南2026年的目标呢，基准还是2025年，增长率取该地区Excel？', session_id='persist')
    assert result['status'] == 'ok'
    assert result['result']['results']['calculate']['value'] == pytest.approx(22992 * 1.10)
    assert result['context_resolution']['document_versions'] == result['state']['fusion_context']['documents']


def test_goal_change_cannot_be_erased_as_reference_syntax(agent):
    assert agent.query(FIRST, session_id='meaning')['status'] == 'ok'
    calls = len(agent.client.contexts)
    result = agent.query('那华南的目标呢', session_id='meaning')
    assert result['status'] == 'clarification'
    assert len(agent.client.contexts) == calls


@pytest.mark.parametrize('question,expected', [
    ('2024年华北地区销售额呢', 3999),
    ('那2025年华南销售额呢', 22992),
    ('那么2025年各地区销售额呢', None),
])
def test_complete_sql_after_fusion_does_not_inherit_document_scope(agent, question, expected):
    assert agent.query(FORECAST, session_id='fresh')['status'] == 'ok'
    def sql_planner(instructions, context, *args, **kwargs):
        assert context['history'] == []
        assert context['question'] == question
        # Model rewrite is deliberately wrong; the original fresh scope wins.
        return {'route': 'sql', 'effective_question': '2026年华东销售额', 'tasks_json': '[]', 'clarification': ''}
    agent.client.generate = sql_planner
    result = agent.query(question, session_id='fresh')
    assert result['status'] == 'ok', result
    assert result['route'] == 'sql'
    assert result['effective_question'] == question
    assert result['context_resolution']['mode'] == 'independent'
    assert result['context_resolution']['reason'] == 'server_verified_self_contained_sql'
    assert not result['state'].get('fusion_context')
    if expected is not None:
        assert result['result']['rows'][0]['销售额'] == expected


@pytest.mark.parametrize('question', [
    '那2025年华南销售额呢，仍按同一公式？',
    '那2025年华南销售额呢，未批准记录不要算',
    '那2025年华南销售额呢，新增未知条件',
])
def test_complete_looking_sql_with_reference_or_unknown_constraint_is_not_fresh(agent, question):
    assert agent.query(FORECAST, session_id='guard')['status'] == 'ok'
    calls = len(agent.client.contexts)
    result = agent.query(question, session_id='guard')
    assert result['status'] == 'clarification'
    assert len(agent.client.contexts) == calls
    assert result['context_resolution'].get('reason') != 'server_verified_self_contained_sql'


def test_true_entity_only_followup_still_inherits(agent):
    assert agent.query(FIRST, session_id='short')['status'] == 'ok'
    result = agent.query('那华南呢', session_id='short')
    assert result['status'] == 'ok'
    assert result['context_resolution']['mode'] == 'server_verified_fusion_followup'
    assert result['result']['results']['calculate']['value'] == pytest.approx(22992 / 3)


def test_complete_sql_can_reset_even_after_failed_fusion(agent):
    assert agent.query(FORECAST, session_id='failedfresh')['status'] == 'ok'
    assert agent.query('那华南呢，未知条件', session_id='failedfresh')['status'] == 'clarification'
    agent.client.generate = lambda *args, **kwargs: {'route': 'sql', 'effective_question': '错', 'tasks_json': '[]', 'clarification': ''}
    result = agent.query('那2025年华南销售额呢', session_id='failedfresh')
    assert result['status'] == 'ok'
    assert result['effective_question'] == '那2025年华南销售额呢'


def test_fresh_sql_does_not_require_old_document_revision(agent):
    assert agent.query(FORECAST, session_id='oldrevision')['status'] == 'ok'
    agent.knowledge.ingest('新版本不含旧公式'.encode(), document_id='policy', title='改版', modality='txt', filename='new.txt')
    def planner(instructions, context, *args, **kwargs):
        assert context['history'] == []
        return {'route': 'sql', 'effective_question': context['question'], 'tasks_json': '[]', 'clarification': ''}
    agent.client.generate = planner
    result = agent.query('2024年华北地区销售额呢', session_id='oldrevision')
    assert result['status'] == 'ok'
    assert result['result']['rows'][0]['销售额'] == 3999


@pytest.mark.parametrize('region,scenario,percent,rate',[('华东','保守',8,.12),('华南','积极',18,.10)])
def test_scenario_switch_compare_summary_use_fresh_source_without_planner(agent,monkeypatch,region,scenario,percent,rate):
    import sqlite3
    workbook=Workbook()
    workbook.active.append(['地区','年份','目标增长率'])
    for entity,growth in [('华东',.12),('华南',.10)]:
        workbook.active.append([entity,2026,growth])
        workbook.active.cell(workbook.active.max_row,3).number_format='0.0%'
    buffer=BytesIO();workbook.save(buffer)
    agent.knowledge.ingest(buffer.getvalue(),document_id='targets',title='区域目标',modality='xlsx',filename='t.xlsx')
    agent.knowledge.ingest(('目标销售额 = 基准销售额 * (1 + 目标增长率)\n'
        '2026年目标增长率为12%，预测基准为2025年销售额。\n'
        '保守目标增长率为8%，积极目标增长率为18%。').encode(),
        document_id='policy',title='指标与预测公式',modality='txt',filename='p.txt')
    question=f'先查2025年{region}销售额作预测基准，再结合《指标与预测公式》的公式和《区域目标》中华东2026目标增长率，算出{region}2026目标销售额'
    question=question.replace('中华东2026',f'中{region}2026')
    first=agent.query(question,session_id='scenario')
    assert first['status']=='ok',first
    monkeypatch.setattr(agent.client,'generate',lambda *a,**kw:pytest.fail('Verified parameter operation called the planner'))
    with sqlite3.connect(agent.engine.database_path) as con:
        base=con.execute("SELECT SUM(sales_amount) FROM sales_orders WHERE region=? AND order_date>='2025-01-01' AND order_date<'2026-01-01'",(region,)).fetchone()[0]
    switched=agent.query(f'基准年份和地区不变，把增长率改用预测公式的{scenario}情景{percent}%，重算{region}目标。',session_id='scenario')
    assert switched['status']=='ok',switched
    assert switched['result']['results']['scenario_target']['value']==pytest.approx(base*(1+percent/100))
    assert switched['result']['results']['scenario_target']['result_unit']=='CNY'
    compared=agent.query(f'{scenario}情景目标比区域目标口径高多少金额？保持相同的2025年{region}销售额基准。',session_id='scenario')
    assert compared['status']=='ok',compared['result'].get('error') or compared['result']
    assert compared['result']['results']['scenario_compare']['difference']==pytest.approx(base*(percent/100-rate))
    summary=agent.query(f'总结{region}2026目标额：区域目标{rate*100:g}%口径和报告{scenario}情景{percent}%口径各是多少？请注明两者都是预测目标，不能当作2026实际销售额。',session_id='scenario')
    assert summary['status']=='ok',summary
    assert '不能当作实际销售额' in summary['result']['answer']
    assert '差额' in summary['result']['answer']
    for key in ('calculate','scenario_target'):
        assert '不能当作实际' in summary['result']['results'][key]['result_interpretation']


@pytest.mark.parametrize('mutation',['percent','origin','revision','year'])
def test_scenario_operation_refuses_unverified_input_before_execution(agent,monkeypatch,mutation):
    agent.knowledge.ingest(('目标销售额 = 基准销售额 * (1 + 目标增长率)\n'
        '2026年目标增长率为12%，预测基准为2025年销售额。\n积极目标增长率为18%。').encode(),
        document_id='policy',title='指标与预测公式',modality='txt',filename='p.txt')
    assert agent.query(FORECAST,session_id='invalid-scenario')['status']=='ok'
    q='基准年份和地区不变，把增长率改用预测公式的积极情景18%，重算华东目标。'
    if mutation=='percent':q=q.replace('18%','19%')
    elif mutation=='year':q=q.replace('重算华东目标','重算华东2027目标')
    elif mutation=='revision':
        agent.knowledge.ingest('积极目标增长率为20%。'.encode(),document_id='policy',title='指标与预测公式',modality='txt',filename='p.txt')
    else:
        turn=agent.conversations.context('invalid-scenario')[-1]
        state=turn.state
        state['fusion_context']['parameter_origin']=state['fusion_context']['task_template']
        state['fusion_context']['parameter_origin']['payload']['tasks'][1]['args']['question']='2024年华南销售额'
        agent.conversations.remember('invalid-scenario',question=turn.question,effective_question=turn.effective_question,state=state)
    monkeypatch.setattr(DependencyAgent,'run',lambda *a,**kw:pytest.fail('Unverified scenario reached execution'))
    result=agent.query(q,session_id='invalid-scenario')
    assert result['status']=='clarification',result


@pytest.mark.parametrize('region,rate',[('华东',.12),('华南',.10)])
def test_explicit_formula_plan_with_unique_real_cell_never_calls_model(agent,monkeypatch,region,rate):
    monkeypatch.setenv('ICT8_FAST_SQL','1')
    monkeypatch.setattr(agent.client,'generate',lambda *a,**kw:pytest.fail('Explicit verified formula request called model'))
    q=f'先统计2025年{region}销售额作为预测基准，结合《指标与预测公式》的公式和《区域目标》中{region}2026目标增长率，计算{region}2026目标销售额。'
    result=agent.query(q,session_id='explicit')
    assert result['status']=='ok',result
    assert result['planner_source']=='server_verified_explicit_formula_plan'
    assert result['result']['results']['target']['value']==pytest.approx((29584 if region=='华东' else 22992)*(1+rate))


@pytest.mark.parametrize('mutation',['unknown_condition','unbound_target','extra_operation','wrong_source'])
def test_explicit_formula_compiler_does_not_drop_unknown_request(agent,mutation):
    from backend.fusion_formula_planner import plan_explicit_formula_request
    q='先查2025年华东销售额作预测基准，再结合《指标与预测公式》的公式和《区域目标》中华东2026目标增长率，算出华东2026目标销售额。'
    if mutation=='unknown_condition':q=q.replace('2025年华东销售额','2025年华东销售额且未批准的不算')
    elif mutation=='unbound_target':q=q.replace('算出华东','算出华北')
    elif mutation=='extra_operation':q=q.rstrip('。')+'，再算实际增长额。'
    else:q=q.replace('《区域目标》','《不存在的表》')
    assert plan_explicit_formula_request(q,agent.engine,agent.knowledge) is None
