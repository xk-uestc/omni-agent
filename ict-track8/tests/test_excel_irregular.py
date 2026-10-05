"""Independent known-answer layout, provenance and ingestion regression tests."""
import base64
from collections import Counter
from io import BytesIO
from pathlib import Path
import sys
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils.cell import get_column_letter, column_index_from_string
from openpyxl.formula.translate import Translator
import pytest

sys.path.insert(0,str(Path(__file__).resolve().parents[2]/'tools'))
from excel_irregular_cases import build_cases,evaluate_case,ExcelCase
from backend.chunk_cleaning import DocumentChunker
from backend.knowledge_store import KnowledgeStore
from backend.dependency_agent import DependencyAgent
from backend.nl2sql.engine import Nl2SqlEngine
from backend.nl2sql.seed import initialize_database
from backend import app as app_module

CASES=build_cases()


def serialize(wb):
    out=BytesIO();wb.save(out);wb.close();return out.getvalue()


def verify_visible_cells(raw,payload):
    """Independent raw-cell oracle, including headers/notes/formula expressions."""
    wb=load_workbook(BytesIO(raw),data_only=False)
    expected={}
    for ws in wb:
        if ws.sheet_state!='visible':continue
        hidden_cols=set()
        for letter,dimension in ws.column_dimensions.items():
            if dimension.hidden:
                start=dimension.min or column_index_from_string(letter);end=dimension.max or start
                hidden_cols.update(range(start,end+1))
        for cell in ws._cells.values():
            if cell.value is not None and not ws.row_dimensions[cell.row].hidden and cell.column not in hidden_cols:
                value=cell.value.isoformat() if hasattr(cell.value,'isoformat') else cell.value
                expected[(ws.title,cell.coordinate)]=(type(value).__name__,value)
    wb.close();observed={};counts=Counter()
    for chunk in payload['chunks']:
        for cell in chunk['metadata'].get('source_cells',[]):
            key=(cell['sheet_name'],cell['coordinate']);value=cell['raw_value']
            actual=(type(value).__name__,value)
            assert key not in observed or observed[key]==actual
            observed[key]=actual;counts[key]+=1
    assert observed==expected
    assert payload['stats']['preserved_visible_cells']==len(expected)
    return counts


@pytest.mark.parametrize('case',CASES,ids=lambda c:c.name)
def test_known_layout_answers_and_visible_cell_conservation(case):
    payload=DocumentChunker().parse_xlsx(case.raw,document_id=case.name).to_dict()
    report=evaluate_case(case,payload)
    assert report['passed'],report
    verify_visible_cells(case.raw,payload)
    if case.name=='14_formula_stale_cache':
        formulas=[v for c in payload['chunks'] for v in c['metadata'].get('values',[]) if isinstance(v,dict) and v.get('formula')]
        assert formulas[0]['cached_value']==999
        assert formulas[0]['cache_status']=='present_unverified'
    if case.name=='16_subtotals':
        assert [c['row_start'] for c in payload['chunks'] if c['metadata'].get('row_role')=='summary']==[3,5]


def translated_case(case,dr,dc):
    source=load_workbook(BytesIO(case.raw));wb=Workbook();wb.remove(wb.active)
    for original in source:
        ws=wb.create_sheet(original.title);ws.sheet_state=original.sheet_state
        for cell in original._cells.values():
            if cell.value is None:continue
            target=ws.cell(cell.row+dr,cell.column+dc)
            target.value=Translator(cell.value,origin=cell.coordinate).translate_formula(target.coordinate) if cell.data_type=='f' else cell.value
            target.number_format=cell.number_format
            if cell.data_type=='e':target.data_type='e'
        for m in original.merged_cells.ranges:
            ws.merge_cells(start_row=m.min_row+dr,end_row=m.max_row+dr,start_column=m.min_col+dc,end_column=m.max_col+dc)
        for r,d in original.row_dimensions.items():
            if d.hidden:ws.row_dimensions[r+dr].hidden=True
        for c,d in original.column_dimensions.items():
            if d.hidden:ws.column_dimensions[get_column_letter(column_index_from_string(c)+dc)].hidden=True
        from copy import deepcopy
        from openpyxl.utils.cell import range_boundaries
        for table in original.tables.values():
            new=deepcopy(table);a,b,c,d=range_boundaries(table.ref)
            new.ref=f'{get_column_letter(a+dc)}{b+dr}:{get_column_letter(c+dc)}{d+dr}';ws.add_table(new)
    source.close();expected=[]
    for record in case.records:
        value=dict(record);value['row']+=dr
        value['headers']=[('column_'+get_column_letter(column_index_from_string(h[7:])+dc)) if h.startswith('column_') else h for h in record['headers']]
        value['values']=[Translator(v,origin=f'A{record["row"]}').translate_formula(f'{get_column_letter(1+dc)}{record["row"]+dr}') if isinstance(v,str) and v.startswith('=') else v for v in record['values']]
        expected.append(value)
    return ExcelCase(case.name,case.description,serialize(wb),expected,case.required_text,case.forbidden_text)


@pytest.mark.parametrize('case',[c for c in CASES if c.name!='18_sparse_far_cell'],ids=lambda c:c.name)
@pytest.mark.parametrize('dr,dc',[(3,2),(11,5),(27,9)])
def test_layout_translation_does_not_change_field_value_relations(case,dr,dc):
    shifted=translated_case(case,dr,dc)
    payload=DocumentChunker().parse_xlsx(shifted.raw,document_id=case.name).to_dict()
    assert evaluate_case(shifted,payload)['passed']
    verify_visible_cells(shifted.raw,payload)


@pytest.mark.parametrize('header_count',[0,1,2])
def test_explicit_unknown_header_or_no_header(header_count):
    wb=Workbook();ws=wb.active;ws.title='未知字段';ws.append(['甲','乙']);ws.append(['上','下']);ws.append(['a',7])
    raw=serialize(wb);payload=DocumentChunker().parse_xlsx(raw,excel_tables=[{'sheet_name':'未知字段','range':'A1:B3','header_rows':header_count}]).to_dict()
    rows=[c for c in payload['chunks'] if c['content_type']=='row']
    assert len(rows)==3-header_count
    assert rows[0]['row_start']==header_count+1
    assert rows[-1]['metadata']['headers']==(['column_A','column_B'] if header_count==0 else ['甲','乙'] if header_count==1 else ['甲 / 上','乙 / 下'])
    verify_visible_cells(raw,payload)


@pytest.mark.parametrize('options',[
    [{'sheet_name':'missing','range':'A1:B2','header_rows':1}],
    [{'sheet_name':'Sheet','range':'B2:A1','header_rows':1}],
    [{'sheet_name':'Sheet','range':'A1:XFE2','header_rows':1}],
    [{'sheet_name':'Sheet','range':'A0:B2','header_rows':1}],
    [{'sheet_name':'Sheet','range':'A1:B2','header_rows':True}],
    [{'sheet_name':'Sheet','range':'A1:B2','header_rows':3}],
    [{'sheet_name':'Sheet','range':'A1:B2','header_rows':1},{'sheet_name':'Sheet','range':'B2:C3','header_rows':0}],
    [{'sheet_name':'Sheet','range':'D1:E3','header_rows':1}],
])
def test_invalid_overlapping_or_empty_region_rejected(options):
    wb=Workbook();wb.active.append(['地区','金额']);wb.active.append(['华东',1])
    with pytest.raises(ValueError):DocumentChunker().parse_xlsx(serialize(wb),excel_tables=options)


def test_merge_budget_checked_before_workbook_loading(monkeypatch):
    wb=Workbook();wb.active['A1']='huge merge';raw=serialize(wb);out=BytesIO()
    with ZipFile(BytesIO(raw)) as src,ZipFile(out,'w',ZIP_DEFLATED) as dst:
        for info in src.infolist():
            content=src.read(info.filename)
            if info.filename=='xl/worksheets/sheet1.xml':
                root=ET.fromstring(content);ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
                node=ET.SubElement(root,ns+'mergeCells',{'count':'1'});ET.SubElement(node,ns+'mergeCell',{'ref':'A1:XFD1048576'})
                content=ET.tostring(root)
            dst.writestr(info,content)
    from backend import excel_layout
    monkeypatch.setattr(excel_layout,'load_workbook',lambda *a,**k:pytest.fail('must reject before allocation'))
    with pytest.raises(ValueError,match='超过上限'):DocumentChunker().parse_xlsx(out.getvalue())


def test_merged_amount_keeps_null_in_covered_rows():
    wb=Workbook();ws=wb.active;ws.append(['地区','金额']);ws.append(['华东',200]);ws.append(['华南',None]);ws.merge_cells('B2:B3');raw=serialize(wb)
    rows=[c for c in DocumentChunker().parse_xlsx(raw).chunks if c.content_type=='row']
    assert rows[1].metadata['values'][1]['raw_value'] is None
    assert rows[1].metadata['values'][1]['merged_anchor']=='B2'
    assert sum(v['raw_value'] or 0 for c in rows for v in [c.metadata['values'][1]])==200


def test_ambiguous_records_that_repeat_are_not_removed_as_headers():
    wb=Workbook();ws=wb.active
    for row in [['华东旗舰','设备甲'],['华南直营','设备乙'],['华东旗舰','设备甲']]:ws.append(row)
    payload=DocumentChunker().parse_xlsx(serialize(wb)).to_dict()
    assert len([c for c in payload['chunks'] if c['content_type']=='row'])==3
    assert any('header_ambiguous' in w for w in payload['warnings'])


def test_wide_long_row_ingestion_retains_sheet_and_cells(tmp_path):
    wb=Workbook();ws=wb.active;ws.title='宽表'
    ws.append(['编号']+[f'字段{i}' for i in range(120)]);ws.append(['unique']+[f'值{i}:'+('长文本'*30) for i in range(120)])
    raw=serialize(wb);store=KnowledgeStore(tmp_path/'knowledge');store.ingest(raw,document_id='wide',title='宽表',modality='xlsx',filename='wide.xlsx')
    document=store.document('wide');verify_visible_cells(raw,{'chunks':document['chunks'],'stats':document['stats']})
    assert all(c['sheet_name']=='宽表' and c['row_start'] is not None for c in document['chunks'])
    assert len([c for c in document['chunks'] if c['content_type']=='row'])>1
    agent=DependencyAgent(Nl2SqlEngine(initialize_database(tmp_path/'db.sqlite')),store)
    result=agent.run([{'id':'cell','tool':'document_cell','args':{'document_id':'wide','where':{'编号':'unique'},'column':'字段119'}}])
    assert result['status']=='ok'
    assert result['results']['cell']['value'].startswith('值119:')


@pytest.mark.parametrize('case',CASES,ids=lambda c:c.name)
def test_preview_endpoint_real_xlsx_upload(case):
    client=TestClient(app_module.app)
    response=client.post('/api/v1/documents/chunks-preview',json={'document_id':case.name,'modality':'xlsx','file_base64':base64.b64encode(case.raw).decode()})
    assert response.status_code==200,response.text
    payload=response.json();assert evaluate_case(case,payload)['passed'];verify_visible_cells(case.raw,payload)


def test_ingest_api_applies_header_confirmation_and_preserves_original(tmp_path,monkeypatch):
    store=KnowledgeStore(tmp_path/'knowledge');monkeypatch.setattr(app_module,'knowledge_store',store)
    wb=Workbook();wb.active.append(['甲','乙']);wb.active.append(['a',7]);raw=serialize(wb)
    payload={'document_id':'confirmed','title':'表头确认','filename':'test.xlsx','modality':'xlsx','file_base64':base64.b64encode(raw).decode(),
             'excel_tables':[{'sheet_name':'Sheet','range':'A1:B2','header_rows':1}]}
    response=TestClient(app_module.app).post('/api/v1/knowledge/ingest',json=payload)
    assert response.status_code==200,response.text
    document=store.document('confirmed');assert document['stats']['table_layouts'][0]['headers']==['甲','乙']
    assert store.original('confirmed')[0].read_bytes()==raw


@pytest.mark.parametrize('name,column,where',[
    ('13_formula_no_cache','金额',{'产品':'设备甲'}),
    ('14_formula_stale_cache','金额',{'产品':'设备甲'}),
    ('17_null_errors','金额',{'地区':'华东'}),
    ('17_null_errors','金额',{'地区':'华南'}),
])
def test_formula_error_and_null_do_not_feed_dependency_calculation(tmp_path,name,column,where):
    case=next(c for c in CASES if c.name==name);store=KnowledgeStore(tmp_path/'knowledge')
    store.ingest(case.raw,document_id='input',title='输入',modality='xlsx',filename='input.xlsx')
    agent=DependencyAgent(Nl2SqlEngine(initialize_database(tmp_path/'db.sqlite')),store)
    result=agent.run([{'id':'cell','tool':'document_cell','args':{'document_id':'input','where':where,'column':column}}])
    assert result['status']!='ok'


def test_merged_dimension_lookup_has_actual_anchor_provenance(tmp_path):
    case=next(c for c in CASES if c.name=='04_merged_dimension');store=KnowledgeStore(tmp_path/'knowledge')
    store.ingest(case.raw,document_id='input',title='输入',modality='xlsx',filename='input.xlsx')
    agent=DependencyAgent(Nl2SqlEngine(initialize_database(tmp_path/'db.sqlite')),store)
    result=agent.run([{'id':'cell','tool':'document_cell','args':{'document_id':'input','where':{'地区':'华东','产品':'设备乙'},'column':'销售额'}}])
    assert result['status']=='ok'
    assert result['results']['cell']['value']==80
    assert result['results']['cell']['condition_coordinates']=={'地区':'A2','产品':'B3'}
    assert result['results']['cell']['cell_coordinate']=='C3'


def test_boolean_selector_cannot_match_integer_zero(tmp_path):
    wb=Workbook();ws=wb.active;ws.append(['启用','金额']);ws.append([False,10]);ws.append([0,20]);raw=serialize(wb)
    store=KnowledgeStore(tmp_path/'knowledge');store.ingest(raw,document_id='input',title='输入',modality='xlsx',filename='input.xlsx')
    agent=DependencyAgent(Nl2SqlEngine(initialize_database(tmp_path/'db.sqlite')),store)
    for flag,wanted in [(False,10),(0,20)]:
        result=agent.run([{'id':'cell','tool':'document_cell','args':{'document_id':'input','where':{'启用':flag},'column':'金额'}}])
        assert result['status']=='ok';assert result['results']['cell']['value']==wanted


def test_unit_row_metadata_flows_to_cell_tool(tmp_path):
    case=next(c for c in CASES if c.name=='24_separate_unit_row');store=KnowledgeStore(tmp_path/'knowledge')
    store.ingest(case.raw,document_id='input',title='输入',modality='xlsx',filename='input.xlsx')
    agent=DependencyAgent(Nl2SqlEngine(initialize_database(tmp_path/'db.sqlite')),store)
    result=agent.run([{'id':'cell','tool':'document_cell','args':{'document_id':'input','where':{'地区':'华东'},'column':'销售额'}}])
    assert result['status']=='ok';assert result['results']['cell']['value']==120;assert result['results']['cell']['unit']=='万元'


def test_user_no_header_overrides_context_note_guess():
    wb=Workbook();ws=wb.active;ws['A1']='说明设备';ws['A2']='备注产品';raw=serialize(wb)
    payload=DocumentChunker().parse_xlsx(raw,excel_tables=[{'sheet_name':'Sheet','range':'A1:A2','header_rows':0}]).to_dict()
    assert len([c for c in payload['chunks'] if c['content_type']=='row'])==2
    verify_visible_cells(raw,payload)


def test_thousands_of_rows_preserve_first_last_and_independent_total():
    wb=Workbook();ws=wb.active;ws.append(['编号','地区','金额'])
    for n in range(1,5001):ws.append([f'ID{n:06d}','华东' if n%2 else '华南',n*3-2])
    raw=serialize(wb);payload=DocumentChunker().parse_xlsx(raw).to_dict()
    rows=[c for c in payload['chunks'] if c['content_type']=='row']
    assert len(rows)==5000;assert rows[0]['metadata']['values'][0]['raw_value']=='ID000001'
    assert rows[-1]['metadata']['values'][0]['raw_value']=='ID005000'
    assert sum(c['metadata']['values'][2]['raw_value'] for c in rows)==3*(5000*5001//2)-2*5000
    verify_visible_cells(raw,payload)
