"""Deterministic, provenance-preserving chunks for office and image documents."""

from __future__ import annotations

import hashlib
import json
import math
import re
import unicodedata
import zipfile
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, time
from io import BytesIO
from pathlib import PurePosixPath
from typing import Any, Iterable, Sequence

from PIL import Image
from pypdf import PdfReader

from .document_analysis import DocumentAnalyzer, PageSignal
from .image_quality import MAX_PIXELS, ImageQualityAnalyzer


_CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_SPACE_RE = re.compile(r"[\t\f\v ]+")
_CJK_RE = re.compile(r"[\u3400-\u9fff]")
_REPEATED_LINE_MIN_PAGES = 2
_DEFAULT_MAX_BYTES = 20 * 1024 * 1024


@dataclass(frozen=True)
class DocumentChunk:
    chunk_id: str
    document_id: str
    modality: str
    content_type: str
    text: str
    title_path: tuple[str, ...] = ()
    page_no: int | None = None
    sheet_name: str | None = None
    row_start: int | None = None
    row_end: int | None = None
    source_locator: str = ""
    parent_chunk_id: str | None = None
    quality: float = 1.0
    warnings: tuple[str, ...] = ()
    metadata: dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        value = asdict(self)
        value["title_path"] = list(self.title_path)
        value["warnings"] = list(self.warnings)
        return value


@dataclass(frozen=True)
class ChunkingResult:
    document_id: str
    modality: str
    chunks: tuple[DocumentChunk, ...]
    stats: dict[str, Any]
    warnings: tuple[str, ...] = ()

    def to_dict(self) -> dict[str, Any]:
        return {
            "document_id": self.document_id,
            "modality": self.modality,
            "chunks": [chunk.to_dict() for chunk in self.chunks],
            "stats": dict(self.stats),
            "warnings": list(self.warnings),
        }


class DocumentChunker:
    """Parse documents into small, stable chunks without losing source evidence."""

    def __init__(
        self,
        max_chars: int = 1800,
        overlap_chars: int = 120,
        *,
        max_bytes: int = _DEFAULT_MAX_BYTES,
        max_excel_cells: int = 500_000,
        max_chunks: int = 100_000,
        max_pdf_ocr_pages: int = 50,
        max_docx_ocr_images: int = 32,
    ):
        if max_chars < 100:
            raise ValueError("max_chars 不能小于 100")
        if overlap_chars < 0 or overlap_chars >= max_chars:
            raise ValueError("overlap_chars 必须大于等于 0 且小于 max_chars")
        if overlap_chars > max_chars // 2:
            raise ValueError("overlap_chars 不能超过 max_chars 的一半")
        if max_bytes < 1 or max_excel_cells < 1 or max_chunks < 1 or max_pdf_ocr_pages < 0 or max_docx_ocr_images < 0:
            raise ValueError("资源上限必须为正数")
        self.max_chars = int(max_chars)
        self.overlap_chars = int(overlap_chars)
        self.max_bytes = int(max_bytes)
        self.max_excel_cells = int(max_excel_cells)
        self.max_chunks = int(max_chunks)
        self.max_pdf_ocr_pages = int(max_pdf_ocr_pages)
        self.max_docx_ocr_images = int(max_docx_ocr_images)

    @staticmethod
    def clean_text(text: str) -> str:
        value = unicodedata.normalize("NFKC", str(text or ""))
        value = value.replace("\r\n", "\n").replace("\r", "\n").replace("\u00a0", " ")
        value = _CONTROL_RE.sub("", value)
        lines = [_SPACE_RE.sub(" ", line).strip() for line in value.split("\n")]
        cleaned: list[str] = []
        for line in lines:
            if line:
                cleaned.append(line)
            elif cleaned and cleaned[-1] != "":
                cleaned.append("")
        while cleaned and cleaned[-1] == "":
            cleaned.pop()
        return "\n".join(cleaned)

    def parse_pdf(
        self,
        pdf_bytes: bytes,
        *,
        document_id: str = "document.pdf",
        ocr_pipeline: Any | None = None,
        language: str = "eng",
    ) -> ChunkingResult:
        self._validate_bytes(pdf_bytes, "PDF")
        try:
            reader = PdfReader(BytesIO(pdf_bytes), strict=False)
            if getattr(reader, "is_encrypted", False):
                raise ValueError("PDF 已加密，无法解析")
            page_count = len(reader.pages)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("PDF 无法解析") from exc
        max_pages = self._positive_env_int("ICT8_PDF_MAX_PAGES", 1000)
        if page_count > max_pages:
            raise ValueError(f"PDF 页数 {page_count} 超过上限 {max_pages}")

        extracted, failed_pages, page_layouts = self._extract_pdf_pages(pdf_bytes, reader)
        page_warnings: dict[int, list[str]] = {
            page_no: ["page_text_extraction_failed"] for page_no in failed_pages
        }

        page_ocr: dict[int, dict[str, Any]] = {}
        min_chars = self._positive_env_int("ICT8_PDF_MIN_TEXT_CHARS", 20)
        ocr_pages = [index for index, text in enumerate(extracted, start=1) if len(re.sub(r"\s", "", text)) < min_chars]
        if len(ocr_pages) > self.max_pdf_ocr_pages:
            deferred_pages = ocr_pages[self.max_pdf_ocr_pages :]
            ocr_pages = ocr_pages[: self.max_pdf_ocr_pages]
            for page_no in deferred_pages:
                page_warnings.setdefault(page_no, []).append("ocr_page_limit_exceeded")
        if ocr_pipeline is not None and ocr_pages:
            try:
                import fitz

                rendered_pdf = fitz.open(stream=pdf_bytes, filetype="pdf")
                for page_no in ocr_pages:
                    render_sha256: str | None = None
                    render_size_px: list[int] | None = None
                    render_scale: float | None = None
                    try:
                        page = rendered_pdf[page_no - 1]
                        rect = page.rect
                        render_scale = min(2.0, 2400 / max(rect.width, rect.height, 1))
                        pixmap = page.get_pixmap(matrix=fitz.Matrix(render_scale, render_scale), alpha=False)
                        if pixmap.width * pixmap.height > 6_000_000:
                            raise ValueError("PDF OCR 页面像素数超过 600 万上限")
                        render_bytes = pixmap.tobytes("png")
                        render_sha256 = hashlib.sha256(render_bytes).hexdigest()
                        render_size_px = [pixmap.width, pixmap.height]
                        if page_no <= len(page_layouts):
                            page_layouts[page_no - 1]["geometry"] = self._pdf_page_geometry(
                                page,
                                render_scale=render_scale,
                                render_size=(pixmap.width, pixmap.height),
                            )
                        result = ocr_pipeline.run(render_bytes, language=language)
                        payload = result.to_dict() if hasattr(result, "to_dict") else dict(result)
                        payload["render_input_sha256"] = render_sha256
                        payload["render_size_px"] = render_size_px
                        payload["render_scale"] = round(render_scale, 6)
                        page_ocr[page_no] = payload
                        extracted[page_no - 1] = str(payload.get("text") or "")
                        page_warnings[page_no] = list(payload.get("warnings") or [])
                        page_warnings[page_no].append(f"ocr_status:{payload.get('status', 'unknown')}")
                    except Exception as exc:
                        page_ocr[page_no] = {
                            "status": "failed",
                            "text": "",
                            "confidence": None,
                            "attempts": [],
                            "selected_transforms": [],
                            "render_input_sha256": render_sha256,
                            "render_size_px": render_size_px,
                            "render_scale": round(render_scale, 6) if render_scale is not None else None,
                        }
                        page_warnings.setdefault(page_no, []).append(f"pdf_page_ocr_failed:{type(exc).__name__}")
                rendered_pdf.close()
            except ImportError:
                for page_no in ocr_pages:
                    page_warnings.setdefault(page_no, []).append("pdf_renderer_unavailable")
            except Exception as exc:
                for page_no in ocr_pages:
                    page_warnings.setdefault(page_no, []).append(f"pdf_renderer_failed:{type(exc).__name__}")
        elif ocr_pages:
            for page_no in ocr_pages:
                page_warnings.setdefault(page_no, []).append("ocr_not_configured")

        ocr_text_limit = 2_000_000
        for page_no, payload in page_ocr.items():
            text_value = str(payload.get("text") or "")
            if len(text_value) > ocr_text_limit:
                page_ocr[page_no] = {"status": "failed", "text": "", "confidence": None, "attempts": payload.get("attempts", []), "selected_transforms": payload.get("selected_transforms", [])}
                extracted[page_no - 1] = ""
                page_warnings.setdefault(page_no, []).append("ocr_text_limit_exceeded")
            elif payload.get("status") != "ok" and text_value.strip():
                page_warnings.setdefault(page_no, []).append("ocr_text_low_confidence_or_incomplete")

        repeated = self._repeated_pdf_lines(extracted)
        removed_line_count = 0
        chunks: list[DocumentChunk] = []
        heading_stack: list[str] = []
        table_headers: list[str] | None = None
        table_title_path: tuple[str, ...] = ()
        last_table_page: int | None = None
        page_signals: list[PageSignal] = []
        for page_no, raw_text in enumerate(extracted, start=1):
            text, removed = self._clean_pdf_page(raw_text, repeated)
            if removed:
                page_lines = raw_text.splitlines()
                removed_line_count += sum(
                    DocumentChunker.clean_text(line) in repeated
                    for index, line in enumerate(page_lines)
                    if index < 2 or index >= max(0, len(page_lines) - 2)
                )
            scanned = len(re.sub(r"\s", "", text)) < self._positive_env_int("ICT8_PDF_MIN_TEXT_CHARS", 20)
            ocr_info = page_ocr.get(page_no, {})
            confidence = self._bounded_quality(ocr_info.get("confidence"))
            if confidence is None:
                confidence = 0.25 if page_no in failed_pages else 0.55 if scanned else 1.0
            page_geometry = page_layouts[page_no - 1].get("geometry", {}) if page_no <= len(page_layouts) else {}
            page_rotation = float(page_geometry.get("rotation_degrees", 0.0) or 0.0)
            page_signals.append(
                PageSignal(
                    page_no,
                    text,
                    ocr_confidence=confidence,
                    rotation_degrees=page_rotation,
                    scanned=scanned,
                )
            )
            outline_warnings: list[str] = []
            page_heading_path, body_blocks = self._pdf_blocks(text, heading_stack, warnings=outline_warnings)
            page_warnings.setdefault(page_no, []).extend(outline_warnings)
            heading_stack = page_heading_path
            for block_index, (block_type, block_text, locator_suffix, block_heading_path) in enumerate(body_blocks):
                locator = f"page:{page_no}:{locator_suffix}"
                page_layout = page_layouts[page_no - 1] if page_no <= len(page_layouts) else {}
                coordinate_evidence = self._pdf_coordinate_evidence(page_layout, block_text)
                block_warnings = list(page_warnings.get(page_no, ()))
                ocr_info = page_ocr.get(page_no, {})
                coordinate_evidence.update(
                    {
                        "render_input_sha256": ocr_info.get("render_input_sha256"),
                        "render_size_px": ocr_info.get("render_size_px"),
                        "render_scale": ocr_info.get("render_scale"),
                    }
                )
                if ocr_info.get("status") == "ok" and not block_warnings:
                    block_warnings.append("ocr_review_recommended")
                if scanned:
                    block_warnings.append("ocr_required")
                if removed:
                    block_warnings.append("repeated_page_matter_removed")
                if block_type == "table":
                    table_rows = self._parse_pdf_table(block_text)
                    if table_headers and table_rows:
                        continuation_allowed = (
                            block_index == 0
                            and last_table_page == page_no - 1
                            and tuple(block_heading_path) == table_title_path
                        )
                        if not continuation_allowed:
                            table_headers = None
                            table_title_path = ()
                        elif len(table_rows[0]) == len(table_headers):
                            normalized_header = [self.clean_text(value).casefold() for value in table_rows[0]]
                            normalized_context = [self.clean_text(value).casefold() for value in table_headers]
                            if normalized_header == normalized_context:
                                table_rows = table_rows[1:]
                            elif self._looks_like_pdf_table_header(table_rows[0]):
                                table_headers = None
                                table_title_path = ()
                            else:
                                block_warnings.append("table_continuation_unverified")
                        if table_headers:
                            table_rows = [table_headers] + table_rows
                            block_heading_path = list(table_title_path or tuple(block_heading_path))
                    if table_rows and not table_headers:
                        table_headers = list(table_rows[0])
                        table_title_path = tuple(block_heading_path)
                    if table_rows:
                        if len(table_rows) == 1 and table_headers and table_rows[0] == table_headers:
                            table_rows = [table_headers]
                        if len(table_rows) >= 1:
                            chunks.extend(
                                self._table_chunks(
                                document_id=document_id,
                                modality="pdf",
                                rows=table_rows,
                                locator_prefix=locator,
                                title_path=tuple(block_heading_path),
                                table_name=f"PDF table on page {page_no}",
                                sheet_name=None,
                                table_metadata={
                                    "page_no": page_no,
                                    "ocr_evidence": self._ocr_evidence(
                                        ocr_info,
                                        language=language,
                                        source_text=raw_text,
                                    ),
                                    "coordinate_evidence": coordinate_evidence,
                                },
                                page_no=page_no,
                                quality=confidence,
                                row_warnings=tuple(dict.fromkeys(block_warnings)),
                                )
                            )
                            last_table_page = page_no
                            self._check_chunk_limit(chunks)
                            continue
                elif block_type == "heading":
                    table_headers = None
                    table_title_path = ()
                    last_table_page = None
                else:
                    table_headers = None
                    table_title_path = ()
                    last_table_page = None
                chunks.extend(
                    self._make_chunks(
                        document_id=document_id,
                        modality="pdf",
                        content_type=block_type,
                        text=block_text,
                        locator=locator,
                        title_path=tuple(block_heading_path),
                        page_no=page_no,
                        quality=confidence,
                        warnings=tuple(dict.fromkeys(block_warnings)),
                        overlap=block_type == "paragraph",
                        metadata={
                            "raw_text": block_text,
                            "raw_page_text": raw_text,
                            "ocr_status": ocr_info.get("status"),
                            "ocr_attempts": ocr_info.get("attempts", []),
                            "ocr_selected_transforms": ocr_info.get("selected_transforms", []),
                            "ocr_metadata": ocr_info.get("metadata", {}),
                            "coordinate_evidence": coordinate_evidence,
                        },
                    )
                )
            self._check_chunk_limit(chunks)
        analysis = DocumentAnalyzer().analyze(document_id=document_id, pages=page_signals)
        warnings = list(f"page_{page}:text_extraction_failed" for page in failed_pages)
        if analysis.ocr_retry_plan.status == "required":
            warnings.append("one_or_more_pages_need_ocr")
        warnings.extend(f"page_{page_no}:{warning}" for page_no, items in page_warnings.items() for warning in items)
        return self._result(
            document_id,
            "pdf",
            chunks,
            warnings,
            extra_stats={
                "page_count": page_count,
                "ocr_required_pages": [item.page_no for item in page_signals if item.scanned],
                "extract_failed_pages": failed_pages,
                "repeated_lines_removed": removed_line_count,
                "document_quality_score": analysis.quality_score,
                "ocr_succeeded_pages": [page_no for page_no, item in page_ocr.items() if item.get("status") == "ok"],
                "ocr_failed_pages": [page_no for page_no, item in page_ocr.items() if item.get("status") != "ok"],
                "rotated_pages": [
                    page_no
                    for page_no, page_layout in enumerate(page_layouts, start=1)
                    if page_layout.get("geometry", {}).get("rotation_degrees", 0) % 360
                ],
                "coordinate_evidence_version": "pdf-coordinate-chain-v1",
            },
        )

    @staticmethod
    def _extract_pdf_pages(pdf_bytes: bytes, reader: Any) -> tuple[list[str], list[int], list[dict[str, Any]]]:
        extracted: list[str] = []
        failed_pages: list[int] = []
        page_layouts: list[dict[str, Any]] = []
        try:
            import fitz

            document = fitz.open(stream=pdf_bytes, filetype="pdf")
            for page_no, page in enumerate(document, start=1):
                try:
                    extracted.append(page.get_text("text", sort=True) or "")
                    page_layouts.append(
                        {
                            "geometry": DocumentChunker._pdf_page_geometry(page),
                            "blocks": DocumentChunker._pdf_text_blocks(page),
                        }
                    )
                except Exception:
                    extracted.append("")
                    page_layouts.append(
                        {
                            "geometry": DocumentChunker._pdf_page_geometry(page),
                            "blocks": [],
                        }
                    )
                    failed_pages.append(page_no)
            document.close()
            return extracted, failed_pages, page_layouts
        except Exception:
            pass
        for page_no, page in enumerate(reader.pages, start=1):
            try:
                extracted.append(page.extract_text() or "")
            except Exception:
                extracted.append("")
                failed_pages.append(page_no)
            page_layouts.append({"geometry": DocumentChunker._pypdf_page_geometry(page), "blocks": []})
        return extracted, failed_pages, page_layouts

    def parse_docx(
        self,
        docx_bytes: bytes,
        *,
        document_id: str = "document.docx",
        ocr_pipeline: Any | None = None,
        language: str = "eng",
    ) -> ChunkingResult:
        self._validate_bytes(docx_bytes, "DOCX")
        self._validate_zip(docx_bytes, "DOCX")
        try:
            from docx import Document
            from docx.oxml.table import CT_Tbl
            from docx.oxml.text.paragraph import CT_P
            from docx.table import Table
            from docx.text.paragraph import Paragraph

            document = Document(BytesIO(docx_bytes))
        except Exception as exc:
            raise ValueError("DOCX 无法解析") from exc

        chunks: list[DocumentChunk] = []
        warnings: list[str] = []
        heading_stack: list[str] = []
        paragraph_no = 0
        table_no = 0
        embedded_image_count = 0
        embedded_image_ocr_count = 0
        embedded_image_ocr_deferred_count = 0
        body = document.element.body
        for element in body.iterchildren():
            if isinstance(element, CT_P):
                paragraph_no += 1
                paragraph = Paragraph(element, document)
                raw = paragraph.text
                math_text = "".join(
                    node.text or ""
                    for node in element.iter()
                    if node.tag.endswith("}t") and "officeDocument/2006/math" in node.tag
                )
                if math_text:
                    raw = f"{raw} {math_text}".strip()
                cleaned = self.clean_text(raw)
                if not cleaned:
                    continue
                style = paragraph.style.name if paragraph.style is not None else ""
                match = re.search(r"(?:heading|标题|überschrift|título)\s*(\d+)", style, re.IGNORECASE)
                outline_level = element.find(
                    ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}outlineLvl"
                )
                is_heading = bool(match) or style.lower() in {"title", "标题"} or outline_level is not None
                if is_heading:
                    if match:
                        level = max(1, int(match.group(1)))
                    elif outline_level is not None:
                        level = max(1, int(outline_level.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", "0")) + 1)
                    else:
                        level = 1
                    heading_stack = heading_stack[: level - 1] + [cleaned]
                    chunks.extend(
                        self._make_chunks(
                            document_id, "docx", "heading", cleaned,
                            locator=f"body:paragraph:{paragraph_no}", title_path=tuple(heading_stack),
                            metadata={"style": style, "raw_text": raw}, overlap=False,
                        )
                    )
                    self._check_chunk_limit(chunks)
                    continue
                num_properties = element.find(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr")
                is_list = num_properties is not None or "list" in style.lower()
                content_type = "list_item" if is_list else "paragraph"
                if is_list:
                    raw_num_id = num_properties.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numId") if num_properties is not None else None
                    raw_level = num_properties.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ilvl") if num_properties is not None else None
                    list_metadata = {
                        "num_id": raw_num_id.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") if raw_num_id is not None else None,
                        "level": raw_level.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") if raw_level is not None else None,
                    }
                else:
                    list_metadata = {}
                chunks.extend(
                    self._make_chunks(
                        document_id, "docx", content_type, cleaned,
                        locator=f"body:paragraph:{paragraph_no}", title_path=tuple(heading_stack),
                        metadata={"style": style, "raw_text": raw, **list_metadata},
                        overlap=not is_list,
                    )
                )
                self._check_chunk_limit(chunks)
            elif isinstance(element, CT_Tbl):
                table_no += 1
                table = Table(element, document)
                rows, has_merged_cells = self._table_rows(table.rows)
                chunks.extend(
                    self._table_chunks(
                        document_id=document_id,
                        modality="docx",
                        rows=rows,
                        locator_prefix=f"body:table:{table_no}",
                        title_path=tuple(heading_stack),
                        table_name=f"Table {table_no}",
                        sheet_name=None,
                        table_metadata={"merged_cells": has_merged_cells},
                    )
                )
                self._check_chunk_limit(chunks)

        for section_no, section in enumerate(document.sections, start=1):
            areas = (
                ("header", section.header),
                ("first_page_header", section.first_page_header),
                ("even_page_header", section.even_page_header),
                ("footer", section.footer),
                ("first_page_footer", section.first_page_footer),
                ("even_page_footer", section.even_page_footer),
            )
            for area_name, area in areas:
                raw = "\n".join(paragraph.text for paragraph in area.paragraphs)
                cleaned = self.clean_text(raw)
                if cleaned:
                    chunks.extend(
                        self._make_chunks(
                            document_id, "docx", area_name, cleaned,
                            locator=f"section:{section_no}:{area_name}", title_path=(), overlap=False,
                            metadata={"raw_text": raw},
                        )
                    )
                for table_no_in_area, table in enumerate(area.tables, start=1):
                    rows, has_merged_cells = self._table_rows(table.rows)
                    chunks.extend(
                        self._table_chunks(
                            document_id=document_id,
                            modality="docx",
                            rows=rows,
                            locator_prefix=f"section:{section_no}:{area_name}:table:{table_no_in_area}",
                            title_path=(),
                            table_name=f"{area_name.title()} table {table_no_in_area}",
                            sheet_name=None,
                            table_metadata={"merged_cells": has_merged_cells, "header_footer": area_name},
                        )
                    )
                self._check_chunk_limit(chunks)

        image_seen: set[str] = set()
        for relation_id, relationship in document.part.rels.items():
            if "image" not in relationship.reltype or relationship.target_ref in image_seen:
                continue
            image_seen.add(relationship.target_ref)
            embedded_image_count += 1
            try:
                image_bytes = relationship.target_part.blob
            except Exception as exc:
                image_bytes = b""
                image_error = f"embedded_image_read_failed:{type(exc).__name__}"
            else:
                image_error = ""
            image_metadata: dict[str, Any] = {
                "relationship_id": relation_id,
                "target": relationship.target_ref,
                "byte_count": len(image_bytes),
                "sha256": hashlib.sha256(image_bytes).hexdigest() if image_bytes else None,
                "ocr_status": "not_run" if ocr_pipeline is None else "pending",
            }
            run_ocr = (
                ocr_pipeline is not None
                and bool(image_bytes)
                and embedded_image_ocr_count < self.max_docx_ocr_images
            )
            image_warnings: list[str] = []
            if run_ocr:
                embedded_image_ocr_count += 1
            elif ocr_pipeline is None:
                image_warnings = ["embedded_image_ocr_not_run"]
            elif not image_bytes:
                image_warnings = ["embedded_image_ocr_not_run"]
            else:
                embedded_image_ocr_deferred_count += 1
                image_warnings = ["embedded_image_ocr_limit_exceeded"]
            if image_error:
                image_warnings.append(image_error)
            ocr_result: dict[str, Any] = {
                "status": "disabled" if ocr_pipeline is None else "failed",
                "text": "",
                "confidence": None,
                "attempts": [],
                "selected_transforms": [],
            }
            if run_ocr:
                try:
                    result = ocr_pipeline.run(image_bytes, language=language)
                    ocr_result = result.to_dict() if hasattr(result, "to_dict") else dict(result)
                except Exception as exc:
                    ocr_result = {
                        **ocr_result,
                        "warnings": [type(exc).__name__],
                        "error_type": type(exc).__name__,
                    }
            image_metadata.update(
                {
                    "ocr_status": ocr_result.get("status"),
                    "ocr_confidence": ocr_result.get("confidence"),
                    "ocr_attempts": ocr_result.get("attempts", []),
                    "ocr_selected_transforms": ocr_result.get("selected_transforms", []),
                    "ocr_executor": ocr_result.get("executor"),
                    "ocr_orientation": ocr_result.get("metadata", {}).get("orientation"),
                }
            )
            image_text = f"Embedded image: {PurePosixPath(relationship.target_ref).name}"
            image_chunk = self._chunk(
                document_id, "docx", "image", f"embedded-image:{relationship.target_ref}",
                image_text,
                metadata=image_metadata,
                warnings=tuple(dict.fromkeys(image_warnings)),
                quality=0.5 if ocr_result.get("status") != "ok" else 0.9,
            )
            chunks.append(image_chunk)
            ocr_text = str(ocr_result.get("text") or "")
            if len(ocr_text) > 2_000_000:
                warnings.append(f"embedded_image_ocr_text_limit_exceeded:{relationship.target_ref}")
                continue
            ocr_warnings = list(ocr_result.get("warnings") or [])
            warnings.extend(ocr_warnings)
            if ocr_result.get("status") != "ok":
                ocr_warnings.append(f"ocr_status:{ocr_result.get('status', 'unknown')}")
            confidence = self._bounded_quality(ocr_result.get("confidence"))
            if confidence is None:
                confidence = 0.5 if ocr_text.strip() else 0.0
            if confidence < 0.8 and ocr_text.strip():
                ocr_warnings.append("ocr_low_confidence")
            for part_no, paragraph in enumerate(self._paragraphs(ocr_text), start=1):
                chunks.extend(
                    self._make_chunks(
                        document_id,
                        "docx",
                        "ocr_region",
                        paragraph,
                        locator=f"embedded-image:{relationship.target_ref}:ocr-region:{part_no}",
                        parent_chunk_id=image_chunk.chunk_id,
                        quality=confidence,
                        warnings=tuple(dict.fromkeys(ocr_warnings)),
                        overlap=False,
                        metadata={
                            "relationship_id": relation_id,
                            "target": relationship.target_ref,
                            "region_no": part_no,
                            "confidence": ocr_result.get("confidence"),
                            "attempts": ocr_result.get("attempts", []),
                            "selected_transforms": ocr_result.get("selected_transforms", []),
                            "ocr_executor": ocr_result.get("executor"),
                            "ocr_metadata": ocr_result.get("metadata", {}),
                            "raw_text": paragraph,
                        },
                    )
                )
            self._check_chunk_limit(chunks)
            self._check_chunk_limit(chunks)
        if not chunks:
            warnings.append("no_extractable_content")
        return self._result(
            document_id,
            "docx",
            chunks,
            warnings,
            extra_stats={
                "paragraph_count": paragraph_no,
                "table_count": table_no,
                "embedded_image_count": embedded_image_count,
                "embedded_image_ocr_count": embedded_image_ocr_count,
                "embedded_image_ocr_deferred_count": embedded_image_ocr_deferred_count,
            },
        )

    def parse_xlsx(self, xlsx_bytes: bytes, *, document_id: str = "document.xlsx") -> ChunkingResult:
        self._validate_bytes(xlsx_bytes, "XLSX")
        self._validate_zip(xlsx_bytes, "XLSX")
        try:
            import openpyxl

            workbook = openpyxl.load_workbook(BytesIO(xlsx_bytes), data_only=False, read_only=False, keep_links=False)
            values_workbook = openpyxl.load_workbook(BytesIO(xlsx_bytes), data_only=True, read_only=True, keep_links=False)
        except Exception as exc:
            raise ValueError("XLSX 无法解析") from exc

        chunks: list[DocumentChunk] = []
        warnings: list[str] = []
        total_cells = 0
        table_count = 0
        for worksheet in workbook.worksheets:
            values_sheet = values_workbook[worksheet.title]
            if worksheet.sheet_state != "visible":
                warnings.append(f"hidden_sheet_skipped:{worksheet.title}")
                continue
            bounds = self._worksheet_bounds(worksheet)
            if bounds is None:
                continue
            max_row, max_column = bounds
            cell_count = max_row * max_column
            total_cells += cell_count
            if total_cells > self.max_excel_cells:
                raise ValueError(f"XLSX 有效单元格超过上限 {self.max_excel_cells}")
            cached_rows: dict[int, tuple[Any, ...]] = {}
            for cached_row_number, cached_row in enumerate(
                values_sheet.iter_rows(
                    min_row=1,
                    max_row=max_row,
                    min_col=1,
                    max_col=max_column,
                    values_only=True,
                ),
                start=1,
            ):
                if any(value is not None for value in cached_row):
                    cached_rows[cached_row_number] = tuple(cached_row)
            hidden_row_numbers = {
                index for index, dimension in worksheet.row_dimensions.items() if dimension.hidden
            }
            hidden_column_indexes = self._hidden_column_indexes(worksheet, max_column)
            hidden_rows = sorted(hidden_row_numbers)
            hidden_columns = [self._column_name(index) for index in sorted(hidden_column_indexes)]
            if hidden_rows:
                warnings.append(f"hidden_rows_skipped:{worksheet.title}:{len(hidden_rows)}")
            if hidden_columns:
                warnings.append(f"hidden_columns_skipped:{worksheet.title}:{len(hidden_columns)}")
            visible_columns = [index for index in range(1, max_column + 1) if index not in hidden_column_indexes]
            raw_rows: list[list[Any]] = []
            for row_index in range(1, max_row + 1):
                if row_index in hidden_row_numbers:
                    continue
                row = [worksheet.cell(row_index, column_index) for column_index in visible_columns]
                values = [cell.value for cell in row]
                if any(value is not None for value in values):
                    cached_row = cached_rows.get(row_index, ())
                    cached_values = [
                        cached_row[column_index - 1] if column_index <= len(cached_row) else None
                        for column_index in visible_columns
                    ]
                    raw_rows.append([row_index, row, values, cached_values])
            if not raw_rows:
                continue
            headers, header_position, header_warning = self._detect_headers(raw_rows)
            if header_warning:
                warnings.append(f"{header_warning}:{worksheet.title}")
            merged_ranges = [str(item) for item in worksheet.merged_cells.ranges]
            hidden_rows = [index for index, dimension in worksheet.row_dimensions.items() if dimension.hidden]
            hidden_columns = [name for name, dimension in worksheet.column_dimensions.items() if dimension.hidden]
            if merged_ranges:
                warnings.append(f"merged_cells:{worksheet.title}:{len(merged_ranges)}")
            table_count += 1
            preserve_header_row = header_position is None or header_warning == "header_ambiguous"
            data_rows = [
                item for index, item in enumerate(raw_rows)
                if preserve_header_row or index != header_position
            ]
            filtered_rows = []
            for row_number, cells, values, cached in data_rows:
                display_values = [self._cell_display(cell, cached[index]) for index, cell in enumerate(cells)]
                if not any(item[0] for item in display_values):
                    continue
                if self._is_repeated_header(values, headers):
                    warnings.append(f"repeated_header_removed:{worksheet.title}:{row_number}")
                    continue
                filtered_rows.append((row_number, cells, display_values))
            if len(chunks) + len(filtered_rows) > self.max_chunks:
                raise ValueError(f"生成 chunk 数超过上限 {self.max_chunks}")
            chunks.extend(
                self._table_chunks(
                    document_id=document_id,
                    modality="xlsx",
                    rows=[headers] + [[self._xlsx_value(display_values[index]) for index in range(len(cells))] for _, cells, display_values in filtered_rows],
                    locator_prefix=f"sheet:{worksheet.title}",
                    title_path=(worksheet.title,),
                    table_name=worksheet.title,
                    sheet_name=worksheet.title,
                    row_numbers=[raw_rows[header_position][0] if header_position is not None else None] + [row_number for row_number, _, _ in filtered_rows],
                    table_metadata={
                        "hidden_sheet": worksheet.sheet_state != "visible",
                        "header_row": raw_rows[header_position][0] if header_position is not None else None,
                        "header_confidence": 0.0 if preserve_header_row else 1.0,
                        "header_decision": "preserved_ambiguous" if preserve_header_row else "selected",
                        "hidden_rows": hidden_rows,
                        "hidden_columns": hidden_columns,
                        "merged_ranges": merged_ranges,
                    },
                )
            )
            self._check_chunk_limit(chunks)
            for row_number, cells, display_values in filtered_rows:
                for column_index, cell in enumerate(cells):
                    if cell.data_type == "f":
                        cache_status = display_values[column_index][1].get("cache_status")
                        warnings.append(
                            f"formula_cache_{cache_status}:{worksheet.title}:{cell.coordinate}"
                        )
                    if cell.data_type == "e":
                        warnings.append(f"cell_error:{worksheet.title}:{cell.coordinate}")
        sheet_count = len(workbook.sheetnames)
        workbook.close()
        values_workbook.close()
        if not chunks:
            warnings.append("no_extractable_content")
        return self._result(document_id, "xlsx", chunks, warnings, extra_stats={"sheet_count": sheet_count, "table_count": table_count, "effective_cells": total_cells})

    def parse_image(
        self,
        image_bytes: bytes,
        *,
        document_id: str = "image",
        ocr_pipeline: Any | None = None,
        language: str = "eng",
    ) -> ChunkingResult:
        self._validate_bytes(image_bytes, "图片")
        try:
            with Image.open(BytesIO(image_bytes)) as image:
                image.verify()
            with Image.open(BytesIO(image_bytes)) as image:
                width, height = image.size
                image_format = image.format
                mode = image.mode
                if width * height > MAX_PIXELS:
                    raise ValueError("图片像素数超过 4000 万上限")
        except Exception as exc:
            if isinstance(exc, ValueError) and "像素数" in str(exc):
                raise
            raise ValueError("图片无法解析") from exc
        try:
            quality_report = ImageQualityAnalyzer().analyze(image_bytes).to_dict()
        except Exception:
            quality_report = {"status": "unavailable"}
        quality_score = self._bounded_quality(quality_report.get("quality_score", quality_report.get("score", 1.0)))
        image_sha256 = hashlib.sha256(image_bytes).hexdigest()
        image_chunk = self._chunk(
            document_id, "image", "image", "image:1", "Image resource; no generated description.",
            identity_context=image_sha256,
            metadata={
                "width": width,
                "height": height,
                "format": image_format,
                "mode": mode,
                "sha256": image_sha256,
                "image_quality": quality_report,
            },
            warnings=("image_caption_unavailable",), quality=quality_score,
        )
        chunks = [image_chunk]
        warnings: list[str] = []
        ocr_result: dict[str, Any] = {"status": "disabled", "text": "", "confidence": None, "attempts": [], "selected_transforms": []}
        if ocr_pipeline is not None:
            try:
                result = ocr_pipeline.run(image_bytes, language=language)
                ocr_result = result.to_dict() if hasattr(result, "to_dict") else dict(result)
            except Exception as exc:
                ocr_result = {"status": "failed", "text": "", "confidence": None, "attempts": [], "selected_transforms": [], "warnings": [type(exc).__name__]}
        else:
            warnings.append("ocr_not_configured")
        ocr_text = str(ocr_result.get("text") or "")
        if len(ocr_text) > 2_000_000:
            raise ValueError("OCR 文本超过 2MB 字符上限")
        confidence = self._bounded_quality(ocr_result.get("confidence"))
        ocr_warnings = list(ocr_result.get("warnings") or [])
        if ocr_result.get("status") != "ok":
            ocr_warnings.append(f"ocr_status:{ocr_result.get('status', 'unknown')}")
        if confidence is None:
            confidence = 0.5 if ocr_text else 0.0
        if confidence < 0.8 and ocr_text.strip():
            ocr_warnings.append("ocr_low_confidence")
        if ocr_text.strip():
            for part_no, paragraph in enumerate(self._paragraphs(ocr_text), start=1):
                for chunk in self._make_chunks(
                    document_id, "image", "ocr_region", paragraph,
                    locator=f"image:1:ocr-region:{part_no}", parent_chunk_id=image_chunk.chunk_id,
                    quality=confidence, warnings=tuple(dict.fromkeys(ocr_warnings)), overlap=False,
                    metadata={
                        "region_no": part_no,
                        "bbox": ocr_result.get("bbox"),
                        "confidence": ocr_result.get("confidence"),
                        "attempts": ocr_result.get("attempts", []),
                        "selected_transforms": ocr_result.get("selected_transforms", []),
                        "ocr_executor": ocr_result.get("executor"),
                        "ocr_metadata": ocr_result.get("metadata", {}),
                        "raw_text": paragraph,
                    },
                ):
                    chunks.append(chunk)
        else:
            warnings.append("ocr_returned_no_text")
        warnings.extend(ocr_warnings)
        self._check_chunk_limit(chunks)
        return self._result(
            document_id, "image", chunks, warnings,
            extra_stats={"image_count": 1, "ocr_status": ocr_result.get("status"), "ocr_confidence": ocr_result.get("confidence"), "ocr_attempt_count": len(ocr_result.get("attempts") or []),
                         "ocr_orientation": ocr_result.get("metadata", {}).get("orientation")},
        )

    def _make_chunks(
        self,
        document_id: str,
        modality: str,
        content_type: str,
        text: str,
        *,
        locator: str,
        title_path: tuple[str, ...] = (),
        page_no: int | None = None,
        sheet_name: str | None = None,
        row_start: int | None = None,
        row_end: int | None = None,
        parent_chunk_id: str | None = None,
        quality: float = 1.0,
        warnings: tuple[str, ...] = (),
        overlap: bool = False,
        metadata: dict[str, Any] | None = None,
        identity_context: str | None = None,
    ) -> list[DocumentChunk]:
        normalized = self.clean_text(text)
        if not normalized:
            return []
        pieces = self._split(normalized, overlap=overlap)
        chunks = []
        for piece_no, piece in enumerate(pieces, start=1):
            piece_locator = locator if len(pieces) == 1 else f"{locator}:part:{piece_no}"
            piece_metadata = dict(metadata or {})
            if len(pieces) > 1:
                piece_metadata["part_no"] = piece_no
                piece_metadata["part_count"] = len(pieces)
            chunks.append(
                self._chunk(
                    document_id, modality, content_type, piece_locator, piece,
                    title_path=title_path, page_no=page_no, sheet_name=sheet_name,
                    row_start=row_start, row_end=row_end, parent_chunk_id=parent_chunk_id,
                    quality=quality, warnings=warnings, metadata=piece_metadata,
                    identity_context=identity_context,
                )
            )
        return chunks

    def _table_chunks(
        self,
        *,
        document_id: str,
        modality: str,
        rows: Sequence[Sequence[Any]],
        locator_prefix: str,
        title_path: tuple[str, ...],
        table_name: str,
        sheet_name: str | None,
        row_numbers: Sequence[int | None] | None = None,
        table_metadata: dict[str, Any] | None = None,
        page_no: int | None = None,
        quality: float = 1.0,
        row_warnings: tuple[str, ...] = (),
    ) -> list[DocumentChunk]:
        if len(rows) < 2:
            if not rows:
                return []
            header_text = self.clean_text(f"{table_name} | " + " | ".join(str(item or f"column_{index + 1}") for index, item in enumerate(rows[0])))
            source_row = row_numbers[0] if row_numbers else 1
            return [
                self._chunk(
                    document_id, modality, "table_header",
                    f"{locator_prefix}:row:{source_row if source_row is not None else 1}",
                    header_text, title_path=title_path, sheet_name=sheet_name,
                    row_start=source_row, row_end=source_row, page_no=page_no, quality=quality,
                    metadata={"table_name": table_name, "headers": list(rows[0]), **(table_metadata or {})},
                    warnings=row_warnings,
                )
            ]
        headers = [str(item or f"column_{index + 1}") for index, item in enumerate(rows[0])]
        results: list[DocumentChunk] = []
        for offset, row in enumerate(rows[1:], start=1):
            values = [self._json_safe(item) for item in row]
            populated = [(headers[index] if index < len(headers) else f"column_{index + 1}", value) for index, value in enumerate(values) if value is not None and value != ""]
            if not populated:
                continue
            source_text = " | ".join(f"{key}: {self._display_value(value)}" for key, value in populated)
            text = self.clean_text(f"{table_name} | " + " | ".join(headers) + "\n" + source_text)
            source_row = row_numbers[offset] if row_numbers and offset < len(row_numbers) else offset + 1
            locator = f"{locator_prefix}:row:{source_row if source_row is not None else offset + 1}"
            parts = self._split(text, overlap=False)
            for part_no, part in enumerate(parts, start=1):
                part_locator = locator if len(parts) == 1 else f"{locator}:part:{part_no}"
                results.append(
                    self._chunk(
                        document_id, modality, "row", part_locator, part, title_path=title_path,
                        sheet_name=sheet_name, row_start=source_row, row_end=source_row, page_no=page_no, quality=quality,
                        metadata={
                            "table_name": table_name,
                            "headers": headers,
                            "values": values,
                            "part_no": part_no if len(parts) > 1 else None,
                            "part_count": len(parts),
                            **(table_metadata or {}),
                        },
                        warnings=tuple(dict.fromkeys(row_warnings + (("merged_cells_present",) if table_metadata and table_metadata.get("merged_ranges") else ()))),
                    )
                )
                self._check_chunk_limit(results)
        return results

    def _chunk(
        self,
        document_id: str,
        modality: str,
        content_type: str,
        locator: str,
        text: str,
        *,
        title_path: tuple[str, ...] = (),
        page_no: int | None = None,
        sheet_name: str | None = None,
        row_start: int | None = None,
        row_end: int | None = None,
        parent_chunk_id: str | None = None,
        quality: float = 1.0,
        warnings: tuple[str, ...] = (),
        metadata: dict[str, Any] | None = None,
        identity_context: str | None = None,
    ) -> DocumentChunk:
        canonical = self.clean_text(text)
        identity = json.dumps(
            [document_id, modality, content_type, locator, canonical, parent_chunk_id, identity_context],
            ensure_ascii=False,
            separators=(",", ":"),
        )
        chunk_id = hashlib.sha256(identity.encode("utf-8")).hexdigest()[:24]
        return DocumentChunk(
            chunk_id=chunk_id,
            document_id=document_id,
            modality=modality,
            content_type=content_type,
            text=canonical,
            title_path=tuple(title_path),
            page_no=page_no,
            sheet_name=sheet_name,
            row_start=row_start,
            row_end=row_end,
            source_locator=locator,
            parent_chunk_id=parent_chunk_id,
            quality=self._bounded_quality(quality) or 0.0,
            warnings=tuple(dict.fromkeys(warnings)),
            metadata=metadata or {},
        )

    def _check_chunk_limit(self, chunks: Sequence[DocumentChunk]) -> None:
        if len(chunks) > self.max_chunks:
            raise ValueError(f"生成 chunk 数超过上限 {self.max_chunks}")

    def _result(
        self,
        document_id: str,
        modality: str,
        chunks: Sequence[DocumentChunk],
        warnings: Sequence[str],
        *,
        extra_stats: dict[str, Any] | None = None,
    ) -> ChunkingResult:
        unique: list[DocumentChunk] = []
        seen: set[tuple[str, str, str]] = set()
        duplicate_count = 0
        for chunk in chunks:
            key = (chunk.modality, chunk.content_type, chunk.text)
            if key in seen:
                duplicate_count += 1
            seen.add(key)
            unique.append(chunk)
        original_chars = sum(len(str(chunk.metadata.get("raw_text", chunk.text))) for chunk in unique)
        cleaned_chars = sum(len(chunk.text) for chunk in unique)
        stats = {
            "raw_character_count": original_chars,
            "cleaned_character_count": cleaned_chars,
            "chunk_count": len(unique),
            "duplicate_content_count": duplicate_count,
            "table_row_count": sum(chunk.content_type == "row" for chunk in unique),
            "formula_count": sum(
                bool(re.search(r"[=＝]", str(chunk.metadata.get("raw_text", chunk.text))))
                or any(isinstance(value, dict) and value.get("formula") for value in chunk.metadata.get("values", []))
                for chunk in unique
            ),
            "low_quality_chunk_count": sum(chunk.quality < 0.8 for chunk in unique),
            "warning_count": sum(len(chunk.warnings) for chunk in unique) + len(warnings),
            "modality_breakdown": self._count_by(unique, "modality"),
            "content_type_breakdown": self._count_by(unique, "content_type"),
            **(extra_stats or {}),
        }
        return ChunkingResult(document_id, modality, tuple(unique), stats, tuple(dict.fromkeys(warnings)))

    def _split(self, text: str, *, overlap: bool) -> list[str]:
        if len(text) <= self.max_chars:
            return [text]
        pieces: list[str] = []
        start = 0
        while start < len(text):
            end = min(len(text), start + self.max_chars)
            if end < len(text):
                lower_bound = start + max(1, self.max_chars // 2)
                boundaries = [text.rfind(mark, lower_bound, end) for mark in ("\n", "。", "！", "？", "；", ".", "!", "?", ";", " ")]
                boundary = max(boundaries)
                if boundary >= lower_bound:
                    end = boundary + 1
            pieces.append(text[start:end].strip())
            if end >= len(text):
                break
            start = max(start + 1, end - self.overlap_chars) if overlap else end
        return [piece for piece in pieces if piece]

    def _overlap_suffix(self, text: str) -> str:
        suffix = text[-self.overlap_chars :] if self.overlap_chars else ""
        if not suffix:
            return ""
        boundary = max(suffix.rfind("。"), suffix.rfind("；"), suffix.rfind("\n"), suffix.rfind(" "))
        return suffix[boundary + 1 :] if boundary >= 0 else suffix

    @staticmethod
    def _paragraphs(text: str) -> list[str]:
        return [item.strip() for item in re.split(r"\n\s*\n|(?<=[。！？.!?])\s+", text) if item.strip()]

    @staticmethod
    def _pdf_text_blocks(page: Any) -> list[dict[str, Any]]:
        try:
            payload = page.get_text("dict", sort=True)
        except Exception:
            return []
        blocks: list[dict[str, Any]] = []
        for block_index, block in enumerate(payload.get("blocks", [])):
            block_type = int(block.get("type", 0))
            bbox = DocumentChunker._rect_values(block.get("bbox"))
            if not bbox:
                continue
            if block_type == 0:
                lines: list[str] = []
                for line in block.get("lines", []):
                    line_text = "".join(str(span.get("text") or "") for span in line.get("spans", []))
                    if line_text:
                        lines.append(line_text)
                normalized_text = DocumentChunker.clean_text("\n".join(lines))
                kind = "text"
            elif block_type == 1:
                normalized_text = ""
                kind = "image"
            else:
                normalized_text = ""
                kind = f"type_{block_type}"
            blocks.append(
                {
                    "block_id": block_index,
                    "kind": kind,
                    "bbox_fitz_unrotated_pt": bbox,
                    "normalized_text": normalized_text,
                    "text_sha256": hashlib.sha256(normalized_text.encode("utf-8")).hexdigest()
                    if normalized_text
                    else None,
                    "char_count": len(normalized_text),
                }
            )
        return blocks

    @staticmethod
    def _pdf_page_geometry(
        page: Any,
        *,
        render_scale: float | None = None,
        render_size: tuple[int, int] | None = None,
    ) -> dict[str, Any]:
        rotation = int(getattr(page, "rotation", 0) or 0) % 360
        media_box = DocumentChunker._rect_values(getattr(page, "mediabox", None))
        crop_box = DocumentChunker._rect_values(getattr(page, "cropbox", None))
        display_rect = DocumentChunker._rect_values(getattr(page, "rect", None))
        render = None
        if render_scale is not None:
            render = {
                "scale": round(float(render_scale), 6),
                "width_px": int(render_size[0]) if render_size else None,
                "height_px": int(render_size[1]) if render_size else None,
                "rotation_applied": True,
                "alpha": False,
                "format": "png",
            }
        return {
            "rotation_degrees": rotation,
            "rotation_source": "pdf_page_rotate",
            "media_box_pdf_pt": media_box,
            "crop_box_pdf_pt": crop_box,
            "display_rect_fitz_pt": display_rect,
            "display_size_pt": [
                round(display_rect[2] - display_rect[0], 4),
                round(display_rect[3] - display_rect[1], 4),
            ]
            if display_rect
            else None,
            "coordinate_systems": {
                "pdf_user_space": {"origin": "bottom_left", "unit": "pt"},
                "fitz_unrotated_text_space": {"origin": "top_left", "unit": "pt"},
                "fitz_display_space": {"origin": "top_left", "unit": "pt"},
                "render_pixel_space": {"origin": "top_left", "unit": "px"},
            },
            "mappings": {
                "pdf_user_to_fitz_unrotated": DocumentChunker._matrix_values(
                    getattr(page, "transformation_matrix", None)
                ),
                "fitz_unrotated_to_display": DocumentChunker._matrix_values(
                    getattr(page, "rotation_matrix", None)
                ),
                "fitz_display_to_render_pixel": [
                    round(float(render_scale), 6),
                    0.0,
                    0.0,
                    round(float(render_scale), 6),
                    0.0,
                    0.0,
                ]
                if render_scale is not None
                else None,
            },
            "render": render,
            "mapping_status": "complete" if display_rect and media_box and crop_box else "partial",
        }

    @staticmethod
    def _pypdf_page_geometry(page: Any) -> dict[str, Any]:
        media_box = DocumentChunker._rect_values(page.mediabox)
        crop_box = DocumentChunker._rect_values(page.cropbox)
        rotation = int(page.get("/Rotate", 0) or 0) % 360
        width = float(page.cropbox.width)
        height = float(page.cropbox.height)
        display_size = [height, width] if rotation in {90, 270} else [width, height]
        return {
            "rotation_degrees": rotation,
            "rotation_source": "pdf_page_rotate",
            "media_box_pdf_pt": media_box,
            "crop_box_pdf_pt": crop_box,
            "display_rect_fitz_pt": [0.0, 0.0, display_size[0], display_size[1]],
            "display_size_pt": display_size,
            "coordinate_systems": {
                "pdf_user_space": {"origin": "bottom_left", "unit": "pt"},
                "fitz_unrotated_text_space": {"origin": "top_left", "unit": "pt"},
                "fitz_display_space": {"origin": "top_left", "unit": "pt"},
                "render_pixel_space": {"origin": "top_left", "unit": "px"},
            },
            "mappings": None,
            "render": None,
            "mapping_status": "fallback_without_fitz",
        }

    @staticmethod
    def _pdf_coordinate_evidence(page_layout: dict[str, Any], block_text: str) -> dict[str, Any]:
        geometry = dict(page_layout.get("geometry") or {})
        target = DocumentChunker.clean_text(block_text)
        matches: list[dict[str, Any]] = []
        for block in page_layout.get("blocks", []):
            candidate = str(block.get("normalized_text") or "")
            if not target or not candidate:
                continue
            if target == candidate or target in candidate or candidate in target:
                matches.append(block)
        source_boxes = [
            block["bbox_fitz_unrotated_pt"]
            for block in matches
            if block.get("bbox_fitz_unrotated_pt")
        ]
        source_bbox = DocumentChunker._union_rects(source_boxes)
        mappings = geometry.get("mappings") or {}
        pdf_to_fitz_matrix = mappings.get("pdf_user_to_fitz_unrotated")
        fitz_to_pdf_matrix = DocumentChunker._invert_matrix(pdf_to_fitz_matrix)
        source_bbox_pdf = (
            DocumentChunker._transform_bbox(source_bbox, fitz_to_pdf_matrix)
            if source_bbox and fitz_to_pdf_matrix
            else None
        )
        display_matrix = mappings.get("fitz_unrotated_to_display")
        display_bbox = DocumentChunker._transform_bbox(source_bbox, display_matrix) if source_bbox and display_matrix else None
        page_bbox = geometry.get("display_rect_fitz_pt")
        effective_display_bbox = display_bbox or page_bbox
        render = geometry.get("render") or {}
        render_scale = render.get("scale")
        render_bbox = (
            DocumentChunker._scale_bbox(effective_display_bbox, render_scale)
            if effective_display_bbox and render_scale is not None
            else None
        )
        return {
            "bbox_status": "exact_block" if source_bbox else "page_level",
            "layout_block_ids": [block.get("block_id") for block in matches],
            "source_bbox_pdf_user_pt": source_bbox_pdf,
            "source_bbox_fitz_unrotated_pt": source_bbox,
            "page_bbox_pdf_user_pt": geometry.get("crop_box_pdf_pt"),
            "source_bbox_fitz_display_pt": effective_display_bbox,
            "source_bbox_render_px": render_bbox,
            "page_geometry": geometry,
        }

    @staticmethod
    def _rect_values(rect: Any) -> list[float] | None:
        if rect is None:
            return None
        try:
            values = [float(value) for value in rect]
        except (TypeError, ValueError):
            return None
        return [round(value, 4) for value in values[:4]] if len(values) >= 4 else None

    @staticmethod
    def _matrix_values(matrix: Any) -> list[float] | None:
        if matrix is None:
            return None
        try:
            values = [float(getattr(matrix, name)) for name in ("a", "b", "c", "d", "e", "f")]
        except (AttributeError, TypeError, ValueError):
            return None
        return [round(value, 6) for value in values]

    @staticmethod
    def _union_rects(rects: Sequence[Sequence[float]]) -> list[float] | None:
        if not rects:
            return None
        return [
            round(min(rect[0] for rect in rects), 4),
            round(min(rect[1] for rect in rects), 4),
            round(max(rect[2] for rect in rects), 4),
            round(max(rect[3] for rect in rects), 4),
        ]

    @staticmethod
    def _transform_bbox(bbox: Sequence[float], matrix: Sequence[float]) -> list[float] | None:
        if len(bbox) < 4 or len(matrix) < 6:
            return None
        a, b, c, d, e, f = matrix[:6]
        corners = (
            (bbox[0], bbox[1]),
            (bbox[2], bbox[1]),
            (bbox[0], bbox[3]),
            (bbox[2], bbox[3]),
        )
        transformed = [(a * x + c * y + e, b * x + d * y + f) for x, y in corners]
        return [
            round(min(point[0] for point in transformed), 4),
            round(min(point[1] for point in transformed), 4),
            round(max(point[0] for point in transformed), 4),
            round(max(point[1] for point in transformed), 4),
        ]

    @staticmethod
    def _invert_matrix(matrix: Sequence[float] | None) -> list[float] | None:
        if matrix is None or len(matrix) < 6:
            return None
        first, second, third, fourth, fifth, sixth = matrix[:6]
        determinant = first * fourth - third * second
        if abs(determinant) < 1e-12:
            return None
        return [
            fourth / determinant,
            -second / determinant,
            -third / determinant,
            first / determinant,
            (third * sixth - fourth * fifth) / determinant,
            (second * fifth - first * sixth) / determinant,
        ]

    @staticmethod
    def _scale_bbox(bbox: Sequence[float] | None, scale: Any) -> list[float] | None:
        if bbox is None or scale is None:
            return None
        try:
            factor = float(scale)
        except (TypeError, ValueError):
            return None
        return [round(float(value) * factor, 4) for value in bbox[:4]]

    @staticmethod
    def _ocr_evidence(payload: dict[str, Any], *, language: str, source_text: str) -> dict[str, Any]:
        return {
            "status": payload.get("status"),
            "executor": payload.get("executor"),
            "language": language,
            "confidence": payload.get("confidence"),
            "attempts": payload.get("attempts", []),
            "selected_transforms": payload.get("selected_transforms", []),
            "bbox": payload.get("bbox"),
            "bbox_status": "available" if payload.get("bbox") is not None else "unavailable",
            "render_input_sha256": payload.get("render_input_sha256"),
            "render_size_px": payload.get("render_size_px"),
            "render_scale": payload.get("render_scale"),
            "source_text_sha256": hashlib.sha256(source_text.encode("utf-8")).hexdigest(),
            "orientation": payload.get("metadata", {}).get("orientation"),
            "coordinate_frame": payload.get("metadata", {}).get("coordinate_frame"),
        }

    @staticmethod
    def _repeated_pdf_lines(pages: Sequence[str]) -> set[str]:
        occurrences: dict[tuple[str, str], set[int]] = {}
        for page_no, text in enumerate(pages, start=1):
            lines = [DocumentChunker.clean_text(line) for line in text.splitlines()]
            edge_lines = [(line, "top") for line in lines[:2]] + [(line, "bottom") for line in lines[-2:]]
            for line, edge in set(edge_lines):
                if line and len(line) <= 120:
                    occurrences.setdefault((line, edge), set()).add(page_no)
        required = max(_REPEATED_LINE_MIN_PAGES, math.ceil(len(pages) * 0.6))
        return {line for (line, _edge), page_numbers in occurrences.items() if len(page_numbers) >= required}

    @staticmethod
    def _clean_pdf_page(text: str, repeated: set[str]) -> tuple[str, bool]:
        lines = [line.rstrip() for line in text.replace("\r", "\n").split("\n")]
        repeated_positions = set(range(min(2, len(lines)))) | set(range(max(0, len(lines) - 2), len(lines)))
        filtered = [line for index, line in enumerate(lines) if not (index in repeated_positions and DocumentChunker.clean_text(line) in repeated)]
        changed = len(filtered) != len(lines)
        joined: list[str] = []
        for raw_line in filtered:
            if "|" not in raw_line and ("\t" in raw_line or re.search(r"\s{3,}", raw_line)):
                cells = re.split(r"\t+|\s{3,}", raw_line.strip())
                cells = [cell.strip() for cell in cells if cell.strip()]
                if len(cells) > 1:
                    raw_line = " | ".join(cells)
            line = DocumentChunker.clean_text(raw_line)
            if not line:
                if joined and joined[-1] != "":
                    joined.append("")
                continue
            if joined and joined[-1] and DocumentChunker._joins_pdf_line(joined[-1], line):
                previous = joined[-1]
                if previous.endswith("-") and re.match(r"^[a-z]", line):
                    joined[-1] = previous[:-1] + line
                else:
                    joined[-1] = previous + ("" if _CJK_RE.search(previous[-1:]) and _CJK_RE.match(line) else " ") + line
            else:
                joined.append(line)
        return "\n".join(joined).strip(), changed

    @staticmethod
    def _joins_pdf_line(previous: str, current: str) -> bool:
        if DocumentChunker._looks_like_heading(current) or DocumentChunker._looks_like_heading(previous):
            return False
        if re.search(r"[。！？；.!?:：]$", previous):
            return False
        if previous.endswith("-") and re.match(r"^[a-z]", current):
            return True
        return bool(_CJK_RE.search(previous[-1:]) and _CJK_RE.match(current)) or previous.endswith((",", "，", "、"))

    @staticmethod
    def _looks_like_heading(line: str) -> bool:
        return bool(DocumentAnalyzer._headings([line]))

    def _pdf_blocks(self, text: str, inherited: list[str], *, warnings: list[str] | None = None) -> tuple[list[str], list[tuple[str, str, str, list[str]]]]:
        headings = list(inherited)
        # Keep declared levels, not list positions: missing heading levels must
        # not invent parents or retain a previous sibling when moving upwards.
        stack = []
        for title in inherited:
            evidence = DocumentAnalyzer._headings([title])
            level = evidence[0].level if evidence else len(stack)+1
            if evidence and evidence[0].rule=='bracket_heuristic':
                level = 2 if stack else 1
            stack.append((level,title))
        body: list[tuple[str, str, str, list[str]]] = []
        lines = text.splitlines()
        current: list[str] = []
        current_type = "paragraph"
        table_no = 0

        def flush() -> None:
            nonlocal current
            value = self.clean_text("\n".join(current))
            if value:
                block_type = "formula" if re.search(r"\b[A-Za-z\u3400-\u9fff]{1,24}\s*[=＝]\s*[^=]{1,100}", value) else current_type
                body.append((block_type, value, f"block:{len(body) + 1}", list(headings)))
            current = []

        for line in lines:
            value = self.clean_text(line)
            if not value:
                if current and current_type == "table":
                    continue
                flush()
                continue
            evidence = DocumentAnalyzer._headings([value])
            if evidence:
                flush()
                heading = evidence[0]
                level = (2 if stack else 1) if heading.rule=='bracket_heuristic' else heading.level
                if warnings is not None:
                    if stack and level > stack[-1][0]+1:
                        warnings.append(f'outline_level_jump:heading:{len(body)+1}')
                    if heading.rule=='bracket_heuristic':
                        warnings.append(f'heuristic_heading:heading:{len(body)+1}')
                stack = [(old_level,title) for old_level,title in stack if old_level < level] + [(level,value)]
                headings = [title for _,title in stack]
                body.append(("heading", value, f"heading:{len(body) + 1}", list(headings)))
                continue
            is_table = "|" in value and value.count("|") >= 2 or "\t" in value or re.search(r"\s{3,}", value)
            next_type = "table" if is_table else "paragraph"
            if current and next_type != current_type:
                flush()
            current_type = next_type
            if next_type == "table" and not current:
                table_no += 1
            current.append(value)
        flush()
        return headings, body

    @staticmethod
    def _table_rows(rows: Iterable[Any]) -> tuple[list[list[str]], bool]:
        result: list[list[str]] = []
        has_merged_cells = False
        for row in rows:
            values = []
            seen_cells: set[int] = set()
            for cell in row.cells:
                identity = id(cell._tc)
                if identity in seen_cells:
                    values.append("")
                    has_merged_cells = True
                else:
                    seen_cells.add(identity)
                    values.append(DocumentChunker.clean_text(cell.text))
            if not values or not any(values):
                continue
            result.append(values)
        return result, has_merged_cells

    @staticmethod
    def _parse_pdf_table(text: str) -> list[list[str]]:
        parsed = []
        for line in (line.strip() for line in text.splitlines() if line.strip()):
            if "|" in line:
                values = [DocumentChunker.clean_text(value) for value in line.strip("|").split("|")]
            elif "\t" in line:
                values = [DocumentChunker.clean_text(value) for value in line.split("\t")]
            else:
                values = [DocumentChunker.clean_text(value) for value in re.split(r"\s{3,}", line)]
            if len(values) > 1 and not all(re.fullmatch(r":?-{3,}:?", value) for value in values):
                parsed.append(values)
        width = max((len(row) for row in parsed), default=0)
        return [row + [""] * (width - len(row)) for row in parsed]

    @staticmethod
    def _looks_like_pdf_table_header(row: Sequence[str]) -> bool:
        values = [DocumentChunker.clean_text(value) for value in row if DocumentChunker.clean_text(value)]
        if len(values) < 2:
            return False
        text_ratio = sum(bool(re.search(r"[^0-9.%-]", value)) for value in values) / len(values)
        unique_ratio = len({value.casefold() for value in values}) / len(values)
        return text_ratio >= 0.75 and unique_ratio >= 0.75

    @staticmethod
    def _worksheet_bounds(worksheet: Any) -> tuple[int, int] | None:
        occupied = [cell for cell in worksheet._cells.values() if cell.value is not None]
        if not occupied:
            return None
        return max(cell.row for cell in occupied), max(cell.column for cell in occupied)

    @staticmethod
    def _hidden_column_indexes(worksheet: Any, max_column: int) -> set[int]:
        hidden: set[int] = set()
        for column_letter, dimension in worksheet.column_dimensions.items():
            if not dimension.hidden:
                continue
            start = int(dimension.min or 0)
            end = int(dimension.max or 0)
            if not start or not end:
                from openpyxl.utils.cell import column_index_from_string

                start = end = column_index_from_string(column_letter)
            hidden.update(index for index in range(start, min(end, max_column) + 1))
        return hidden

    @staticmethod
    def _column_name(index: int) -> str:
        from openpyxl.utils.cell import get_column_letter

        return get_column_letter(index)

    @staticmethod
    def _detect_headers(raw_rows: Sequence[Any]) -> tuple[list[str], int | None, str | None]:
        candidates = raw_rows[:10]
        scored: list[tuple[float, int, list[Any]]] = []
        for index, (row_number, _cells, values, _cached) in enumerate(candidates):
            nonempty = [value for value in values if value is not None and str(value).strip()]
            if not nonempty:
                continue
            text_ratio = sum(isinstance(value, str) for value in nonempty) / len(nonempty)
            unique_ratio = len({str(value).strip().casefold() for value in nonempty}) / len(nonempty)
            score = len(nonempty) + text_ratio * 0.75 + unique_ratio * 0.25 - index * 0.12
            scored.append((score, index, values))
        if not scored:
            return [], None, "header_ambiguous"
        _score, selected, values = max(scored, key=lambda item: item[0])
        nonempty = [value for value in values if value is not None and str(value).strip()]
        text_ratio = sum(isinstance(value, str) for value in nonempty) / max(1, len(nonempty))
        unique_ratio = len({str(value).strip().casefold() for value in nonempty}) / max(1, len(nonempty))
        has_header = text_ratio >= 0.75 and unique_ratio >= 0.75
        preceding_rows = candidates[:selected]
        if selected > 0 and any(
            sum(value is not None and str(value).strip() for value in row_values) > 1
            for _row_number, _cells, row_values, _cached in preceding_rows
        ):
            has_header = False
        if not has_header:
            return [f"column_{index + 1}" for index in range(len(values))], None, "header_ambiguous"
        names: list[str] = []
        used: dict[str, int] = {}
        for column, value in enumerate(values, start=1):
            base = DocumentChunker.clean_text(str(value)) if value is not None else ""
            base = base or f"column_{column}"
            used[base.casefold()] = used.get(base.casefold(), 0) + 1
            names.append(base if used[base.casefold()] == 1 else f"{base}_{used[base.casefold()]}")
        return names, selected, None

    @staticmethod
    def _is_repeated_header(values: Sequence[Any], headers: Sequence[str]) -> bool:
        normalized = [DocumentChunker.clean_text(str(value)) if value is not None else "" for value in values]
        return len(normalized) == len(headers) and [value.casefold() for value in normalized] == [value.casefold() for value in headers]

    @staticmethod
    def _cell_display(cell: Any, cached_value: Any) -> tuple[bool, dict[str, Any]]:
        value = cell.value
        if value is None:
            return False, {"raw_value": None}
        display = DocumentChunker._json_safe(value)
        metadata: dict[str, Any] = {"raw_value": display, "number_format": cell.number_format}
        if cell.data_type == "f":
            metadata["formula"] = value
            metadata["cached_value"] = DocumentChunker._json_safe(cached_value)
            metadata["cache_status"] = "present_unverified" if cached_value is not None else "missing"
            return True, metadata
        if cell.data_type == "e":
            metadata["error"] = value
        return True, metadata

    @staticmethod
    def _xlsx_value(value: tuple[bool, dict[str, Any]]) -> dict[str, Any] | None:
        return value[1] if value[0] else None

    @staticmethod
    def _display_value(value: Any) -> str:
        if isinstance(value, dict):
            if value.get("formula") is not None:
                rendered = f"formula={value['formula']}"
                if value.get("cached_value") is not None:
                    rendered += f" (cached={value['cached_value']})"
                return rendered
            if "raw_value" in value:
                raw_value = value["raw_value"]
                number_format = str(value.get("number_format") or "")
                if isinstance(raw_value, (int, float)) and "%" in number_format:
                    return f"{raw_value * 100:g}%"
                currency_symbols = ("$", "€", "£", "¥")
                if isinstance(raw_value, (int, float)) and any(symbol in number_format for symbol in currency_symbols):
                    symbol = next(symbol for symbol in currency_symbols if symbol in number_format)
                    decimals = 2 if ".00" in number_format or ".0" in number_format else 0
                    return f"{symbol}{raw_value:,.{decimals}f}"
                if isinstance(raw_value, (datetime, date, time)):
                    return raw_value.isoformat()
                return str(raw_value)
            return json.dumps(value, ensure_ascii=False, sort_keys=True)
        if isinstance(value, (tuple, list)):
            return json.dumps(value, ensure_ascii=False, sort_keys=True)
        return str(value)

    @staticmethod
    def _json_safe(value: Any) -> Any:
        if isinstance(value, (datetime, date, time)):
            return value.isoformat()
        if isinstance(value, float) and not math.isfinite(value):
            return str(value)
        if value is None or isinstance(value, (str, int, float, bool)):
            return value
        if isinstance(value, dict):
            return {str(key): DocumentChunker._json_safe(item) for key, item in value.items()}
        if isinstance(value, (list, tuple)):
            return [DocumentChunker._json_safe(item) for item in value]
        return str(value)

    @staticmethod
    def _count_by(chunks: Sequence[DocumentChunk], field_name: str) -> dict[str, int]:
        counts: dict[str, int] = {}
        for chunk in chunks:
            key = str(getattr(chunk, field_name))
            counts[key] = counts.get(key, 0) + 1
        return counts

    @staticmethod
    def _bounded_quality(value: Any) -> float | None:
        if value is None:
            return None
        try:
            parsed = float(value)
        except (TypeError, ValueError):
            return None
        if not math.isfinite(parsed):
            return None
        if parsed > 1.0 and parsed <= 100.0:
            parsed /= 100.0
        return round(max(0.0, min(1.0, parsed)), 4)

    @staticmethod
    def _positive_env_int(name: str, default: int) -> int:
        import os

        try:
            value = int(os.getenv(name, str(default)) or default)
            return value if value > 0 else default
        except (TypeError, ValueError):
            return default

    def _validate_bytes(self, raw: bytes, file_type: str) -> None:
        if not isinstance(raw, bytes) or not raw:
            raise ValueError(f"{file_type} 不能为空")
        if len(raw) > self.max_bytes:
            raise ValueError(f"{file_type} 超过 {self.max_bytes} 字节限制")

    @staticmethod
    def _validate_zip(raw: bytes, file_type: str) -> None:
        try:
            with zipfile.ZipFile(BytesIO(raw)) as archive:
                infos = archive.infolist()
                expanded_size = sum(info.file_size for info in infos)
                if expanded_size > 100 * 1024 * 1024:
                    raise ValueError(f"{file_type} 解压后体积超过 100MB 限制")
                if len(infos) > 20_000:
                    raise ValueError(f"{file_type} 内部文件数量超过上限")
                for info in infos:
                    if info.file_size > 25 * 1024 * 1024:
                        raise ValueError(f"{file_type} 单个内部文件超过 25MB 限制")
                    if info.file_size > 1 * 1024 * 1024 and info.compress_size and info.file_size / info.compress_size > 200:
                        raise ValueError(f"{file_type} 压缩比超过安全上限")
        except zipfile.BadZipFile as exc:
            raise ValueError(f"{file_type} 无法解析") from exc
